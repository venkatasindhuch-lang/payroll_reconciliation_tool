import io
import re
import unicodedata
from difflib import SequenceMatcher
from datetime import date
import os
import itertools

import pandas as pd
import streamlit as st

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


try:
    from PIL import Image, ImageOps, ImageEnhance
except ImportError:
    Image = None
    ImageOps = None
    ImageEnhance = None

try:
    import pytesseract
    from pytesseract import Output as TesseractOutput
except ImportError:
    pytesseract = None
    TesseractOutput = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

st.set_page_config(page_title="CARRR Payroll", page_icon="✅", layout="wide")

# -----------------------------
# CARRR Payroll configuration
# -----------------------------
RECON_FIELDS = [
    "Gross Pay", "Salary", "PAYE", "Employee NI", "Pension",
    "Other Deductions", "Net Pay", "Overtime", "Overtime Pay", "Bonus"
]
REQUIRED_KEYS = ["Employee ID", "Employee Name"]

# CARRR Control Centre styling
st.markdown(
    """
    <style>
    * {box-sizing:border-box}
    .stApp{background:linear-gradient(180deg,#f8fbff 0%,#ffffff 38%,#fbf9ff 100%);color:#17213d}
    .block-container{padding-top:2.1rem;padding-bottom:3rem;max-width:1180px}
    h1,h2,h3{letter-spacing:-.025em}
    .carrr-head{display:flex;justify-content:space-between;gap:1rem;align-items:flex-start;padding:1.15rem 1.3rem;border:1px solid #e9e8f4;border-radius:22px;background:rgba(255,255,255,.92);box-shadow:0 12px 32px rgba(41,52,94,.06);margin-bottom:1rem}
    .carrr-brand{font-size:1.9rem;font-weight:850;color:#141d4c}.carrr-tag{font-size:.95rem;font-weight:750;color:#5a63a8;margin-top:.15rem}.carrr-sub{color:#71758b;font-size:.88rem;margin-top:.25rem}
    .phasebar{display:grid;grid-template-columns:.85fr 1fr .7fr 1fr .8fr;gap:.5rem;margin:.8rem 0 1.2rem}
    .phase{padding:.72rem .8rem;border-radius:15px;border:1px solid #e5e6f2;background:#fff;text-align:center;font-size:.78rem;font-weight:800;color:#3d4666}
    .phase.input{background:#fff7ed}.phase.before{background:#edf5ff}.phase.run{background:#f6f5f8}.phase.after{background:#f2efff}.phase.gate{background:#effaf4}
    .summary-card{border:1px solid #e8e9f3;background:#fff;border-radius:18px;padding:1rem 1.05rem;min-height:112px}.summary-label{font-size:.75rem;font-weight:800;color:#74788f;text-transform:uppercase;letter-spacing:.06em}.summary-value{font-size:1.75rem;font-weight:850;margin-top:.24rem;color:#17213d}.summary-note{font-size:.8rem;color:#7a7e92;margin-top:.2rem}
    .before-box,.after-box,.gate-box{border-radius:22px;padding:1.2rem 1.25rem;border:1px solid #e7e8f3;background:#fff;margin:.7rem 0 1rem}.before-box{border-left:5px solid #3b82f6}.after-box{border-left:5px solid #7958e8}.gate-box{border-left:5px solid #10a15b}
    .kicker{font-size:.74rem;letter-spacing:.14em;font-weight:850;color:#6e75a6;text-transform:uppercase}.bigline{font-size:1.45rem;font-weight:850;color:#182044;margin:.15rem 0}.muted{color:#73788d;font-size:.9rem}
    .status-pass{color:#087f4a;font-weight:800}.status-warn{color:#a56600;font-weight:800}.status-fail{color:#c72f3c;font-weight:800}
    .exception-head{padding:.9rem 1rem;border-radius:16px;background:#fff7ef;border:1px solid #ffe2bd;margin:.7rem 0}.release-ready{padding:1.1rem 1.2rem;border-radius:18px;background:#edfbf3;border:1px solid #bfead0;color:#086c40;font-size:1.4rem;font-weight:900;text-align:center}.release-hold{padding:1.1rem 1.2rem;border-radius:18px;background:#fff1f2;border:1px solid #ffcdd3;color:#a51f31;font-size:1.4rem;font-weight:900;text-align:center}
    .gate-auto-clear{padding:.85rem 1rem;border-radius:14px;background:#f0fdf4;border:1px solid #bbf7d0;color:#166534;font-weight:750;margin:.4rem 0 .8rem 0}
    .gate-auto-hold{padding:.85rem 1rem;border-radius:14px;background:#fff7ed;border:1px solid #fed7aa;color:#9a3412;font-weight:750;margin:.4rem 0 .8rem 0}
    .privacy{font-size:.78rem;color:#777b8f;padding:.55rem .7rem;border-radius:12px;background:#fffdf4;border:1px solid #f1ead0;margin:.5rem 0 1rem}
    div[data-testid="stMetric"]{background:white;border:1px solid #ececf4;padding:.75rem .85rem;border-radius:16px}
    @media(max-width:800px){.phasebar{grid-template-columns:1fr 1fr}.carrr-head{display:block}.block-container{padding-left:1rem;padding-right:1rem}.carrr-brand{font-size:1.6rem}}
    </style>
    """, unsafe_allow_html=True
)

st.markdown("""
<div class="carrr-head">
 <div>
  <div class="carrr-brand">CARRR Payroll</div>
  <div class="carrr-tag">Before payroll. After payroll. Before release.</div>
  <div class="carrr-sub">Check · Analyse · Reconcile · Review · Release</div>
 </div>
 <div class="muted">Payroll control centre</div>
</div>
<div class="phasebar">
 <div class="phase input">INPUT HUB<br>Capture + Verify</div>
 <div class="phase before">BEFORE PAYROLL<br>Check + Analyse</div>
 <div class="phase run">PAYROLL RUN<br>Your payroll system</div>
 <div class="phase after">AFTER PAYROLL<br>Reconcile + Review</div>
 <div class="phase gate">RELEASE GATE<br>Ready / Hold</div>
</div>
<div class="privacy"><b>Public trial:</b> use fictional or anonymised payroll data. Real payroll data should only be handled in an appropriately secured private deployment.</div>
""", unsafe_allow_html=True)

def clean_columns(df):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    return df



# -----------------------------
# Smart Header Detection + Mapping
# -----------------------------
CARRR_MAPPING_FIELDS = [
    "Employee ID", "Employee Name", "First Name", "Last Name", "Tax Code", "NI Number",
    "Bank Details", "Sort Code", "Account Number", "Employment Status", "Starter/Leaver",
    "Salary", "Overtime Hours", "Overtime Pay", "Bonus", "Gross Pay", "PAYE",
    "Employee NI", "Pension", "Other Deductions", "Net Pay"
]

HEADER_ALIASES = {
    "Employee ID": [
        "employee id", "employee number", "employee no", "employee num", "emp id", "emp no",
        "emp number", "staff id", "staff no", "staff number", "worker id", "worker no",
        "worker number", "payroll id", "payroll no", "payroll number", "person id", "person number",
        "personnel no", "personnel number", "associate id", "associate number", "employee ref",
        "worker ref", "staff ref", "employee code", "staff code"
    ],
    "Employee Name": [
        "employee name", "emp name", "staff name", "worker name", "full name", "name of employee",
        "employee full name", "worker full name", "staff full name", "payee name", "person name"
    ],
    "First Name": ["first name", "forename", "given name", "employee first name", "worker first name"],
    "Last Name": ["last name", "surname", "family name", "employee surname", "worker surname"],
    "Tax Code": ["tax code", "paye code", "hmrc tax code", "employee tax code"],
    "NI Number": [
        "ni number", "ni no", "national insurance number", "national insurance no", "nino",
        "national insurance", "employee ni number"
    ],
    "Bank Details": ["bank details", "bank info", "bank information", "payment bank details"],
    "Sort Code": ["sort code", "bank sort code", "sortcode"],
    "Account Number": ["account number", "account no", "bank account", "bank account number", "acct no", "acct number"],
    "Employment Status": ["employment status", "employee status", "worker status", "status", "payroll status"],
    "Starter/Leaver": [
        "starter leaver", "starter/leaver", "joiner leaver", "joiner/leaver", "starter status",
        "leaver status", "new starter", "leaver", "starter"
    ],
    "Salary": [
        "salary", "basic salary", "base salary", "monthly salary", "basic pay", "base pay",
        "contract salary", "regular pay", "standard pay"
    ],
    "Overtime Hours": [
        "overtime hours", "ot hours", "ot hrs", "overtime hrs", "extra hours", "additional hours",
        "overtime units", "ot units"
    ],
    "Overtime Pay": [
        "overtime pay", "ot pay", "overtime amount", "ot amount", "overtime value", "extra pay",
        "overtime earnings", "ot earnings"
    ],
    "Bonus": ["bonus", "bonus pay", "bonus amount", "commission", "commission pay", "incentive", "incentive pay"],
    "Gross Pay": [
        "gross pay", "gross", "gross amount", "gross earnings", "total gross", "gross salary",
        "total earnings", "gross taxable pay", "gross taxable"
    ],
    "PAYE": [
        "paye", "tax", "income tax", "tax deducted", "tax deduction", "paye tax", "paye deduction",
        "employee tax", "tax amount"
    ],
    "Employee NI": [
        "employee ni", "ni", "national insurance", "ni deduction", "ni deducted", "employee nic",
        "employee national insurance", "nic", "ee ni", "ee nic"
    ],
    "Pension": [
        "pension", "employee pension", "pension deduction", "pension contribution", "employee pension contribution",
        "ee pension", "workplace pension"
    ],
    "Other Deductions": [
        "other deductions", "other deduction", "deductions", "other ded", "misc deductions",
        "miscellaneous deductions", "additional deductions", "total other deductions"
    ],
    "Net Pay": [
        "net pay", "net", "take home", "take home pay", "takehome", "takehome pay", "net amount",
        "net salary", "net payment", "amount paid", "payment amount"
    ],
}


def normalize_header(value):
    text_value = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    text_value = text_value.lower().replace("&", " and ")
    text_value = re.sub(r"[^a-z0-9]+", " ", text_value)
    return " ".join(text_value.split()).strip()


_ALIAS_LOOKUP = {
    normalize_header(alias): field
    for field, aliases in HEADER_ALIASES.items()
    for alias in ([field] + aliases)
}


def _token_score(a, b):
    a_tokens = set(a.split())
    b_tokens = set(b.split())
    if not a_tokens or not b_tokens:
        return 0.0
    return len(a_tokens & b_tokens) / len(a_tokens | b_tokens)


def suggest_carrr_field(header):
    normalized = normalize_header(header)
    if not normalized:
        return None, "Low", 0.0

    if normalized in _ALIAS_LOOKUP:
        return _ALIAS_LOOKUP[normalized], "High", 1.0

    best_field = None
    best_score = 0.0
    for field, aliases in HEADER_ALIASES.items():
        for alias in [field] + aliases:
            alias_norm = normalize_header(alias)
            seq = SequenceMatcher(None, normalized, alias_norm).ratio()
            token = _token_score(normalized, alias_norm)
            containment = 1.0 if (normalized in alias_norm or alias_norm in normalized) and min(len(normalized), len(alias_norm)) >= 5 else 0.0
            score = max((0.65 * seq) + (0.35 * token), 0.88 * containment)
            if score > best_score:
                best_score = score
                best_field = field

    # Small semantic rules for common payroll abbreviations.
    tokens = set(normalized.split())
    if {"emp", "id"}.issubset(tokens) or {"worker", "id"}.issubset(tokens) or {"staff", "id"}.issubset(tokens):
        return "Employee ID", "High", max(best_score, 0.95)
    if "name" in tokens and any(t in tokens for t in ["employee", "emp", "worker", "staff", "full"]):
        return "Employee Name", "High", max(best_score, 0.95)
    if ("ot" in tokens or "overtime" in tokens) and ("hour" in tokens or "hours" in tokens or "hrs" in tokens):
        return "Overtime Hours", "High", max(best_score, 0.95)
    if "net" in tokens and ("pay" in tokens or "amount" in tokens or len(tokens) == 1):
        return "Net Pay", "High", max(best_score, 0.94)
    if "gross" in tokens and ("pay" in tokens or "amount" in tokens or len(tokens) == 1):
        return "Gross Pay", "High", max(best_score, 0.94)

    if best_score >= 0.86:
        confidence = "High"
    elif best_score >= 0.68:
        confidence = "Medium"
    else:
        confidence = "Low"
        best_field = None
    return best_field, confidence, round(best_score, 2)


def make_unique_headers(values):
    seen = {}
    headers = []
    for i, value in enumerate(values):
        base = str(value).strip() if pd.notna(value) else ""
        if not base or base.lower().startswith("unnamed"):
            base = f"Column {i + 1}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base} ({count + 1})")
    return headers


def detect_header_row(raw_df, max_rows=15):
    if raw_df.empty:
        return 0
    best_row = 0
    best_score = -1.0
    rows_to_check = min(max_rows, len(raw_df))
    for row_idx in range(rows_to_check):
        values = raw_df.iloc[row_idx].tolist()
        mapped = 0
        confidence_points = 0.0
        non_empty = 0
        for value in values:
            if pd.isna(value) or str(value).strip() == "":
                continue
            non_empty += 1
            _, confidence, score = suggest_carrr_field(value)
            if confidence == "High":
                mapped += 1
                confidence_points += 2.0 + score
            elif confidence == "Medium":
                mapped += 1
                confidence_points += 1.0 + score
        row_score = (mapped * 3.0) + confidence_points + min(non_empty, 10) * 0.05
        if row_score > best_score:
            best_score = row_score
            best_row = row_idx
    # If no payroll-like headers were found, preserve the conventional first row.
    return best_row if best_score >= 3.0 else 0


def dataframe_from_raw(raw_df):
    header_row = detect_header_row(raw_df)
    headers = make_unique_headers(raw_df.iloc[header_row].tolist())
    df = raw_df.iloc[header_row + 1:].copy()
    df.columns = headers
    df = df.dropna(how="all").reset_index(drop=True)
    # Remove fully blank columns after the header is set.
    blank_cols = [c for c in df.columns if df[c].isna().all()]
    if blank_cols:
        df = df.drop(columns=blank_cols)
    return clean_columns(df), header_row + 1


def apply_mapping(df, mapping):
    selected = {client: target for client, target in mapping.items() if target and target != "Ignore"}
    mapped = df[list(selected.keys())].rename(columns=selected).copy()
    # Convert first/last name inputs into Employee Name where possible.
    if "Employee Name" not in mapped.columns and ("First Name" in mapped.columns or "Last Name" in mapped.columns):
        first = mapped["First Name"].fillna("").astype(str).str.strip() if "First Name" in mapped.columns else ""
        last = mapped["Last Name"].fillna("").astype(str).str.strip() if "Last Name" in mapped.columns else ""
        if isinstance(first, str):
            mapped["Employee Name"] = last
        elif isinstance(last, str):
            mapped["Employee Name"] = first
        else:
            mapped["Employee Name"] = (first + " " + last).str.strip()
    return clean_columns(mapped)



WEEKDAY_PATTERNS = {
    "Mon": ["mon", "monday"],
    "Tue": ["tue", "tues", "tuesday"],
    "Wed": ["wed", "weds", "wednesday"],
    "Thu": ["thu", "thur", "thurs", "thursday"],
    "Fri": ["fri", "friday"],
    "Sat": ["sat", "saturday"],
    "Sun": ["sun", "sunday"],
}


def _find_column_by_aliases(df, aliases):
    normalized = {normalize_header(col): col for col in df.columns}
    for alias in aliases:
        a = normalize_header(alias)
        if a in normalized:
            return normalized[a]
    # conservative containment fallback
    for col_norm, original in normalized.items():
        for alias in aliases:
            a = normalize_header(alias)
            if len(a) >= 5 and (a in col_norm or col_norm in a):
                return original
    return None


def detect_weekday_columns(df):
    found = {}
    for label, aliases in WEEKDAY_PATTERNS.items():
        for col in df.columns:
            norm = normalize_header(col)
            for alias in aliases:
                if re.fullmatch(rf"{re.escape(alias)}( hrs?| hours?| time)?", norm):
                    found[label] = col
                    break
            if label in found:
                break
    return found


def detect_source_type(df):
    fields = set()
    for col in df.columns:
        suggested, confidence, _ = suggest_carrr_field(col)
        if suggested and confidence in ["High", "Medium"]:
            fields.add(suggested)

    weekdays = detect_weekday_columns(df)
    std_col = _find_column_by_aliases(df, ["std hours", "standard hours", "contract hours", "normal hours", "regular hours"])
    rate_col = _find_column_by_aliases(df, ["pay rate", "hourly rate", "hour rate", "rate per hour"])
    factor_col = _find_column_by_aliases(df, ["ot factor", "overtime factor", "ot multiplier", "overtime multiplier"])

    if len(weekdays) >= 3 or (std_col and (rate_col or "Overtime Hours" in fields)):
        return "Timesheet", {"weekdays": weekdays, "std_col": std_col, "rate_col": rate_col, "factor_col": factor_col}

    output_markers = {"Gross Pay", "PAYE", "Employee NI", "Pension", "Net Pay"}
    if len(fields & output_markers) >= 3:
        return "Payroll output / register", {}

    master_markers = {"Tax Code", "NI Number", "Bank Details", "Salary", "Starter/Leaver", "Employment Status"}
    if len(fields & master_markers) >= 2:
        return "Payroll input / changes", {}

    if "Bonus" in fields and len(fields) <= 5:
        return "Bonus / variable pay input", {}
    if "Overtime Hours" in fields or "Overtime Pay" in fields:
        return "Overtime input", {}
    return "Payroll data", {}


def build_auto_mapping(df):
    suggestions = []
    candidates = {}
    for col in df.columns:
        suggested, confidence, score = suggest_carrr_field(col)
        suggestions.append((col, suggested, confidence, score))
        # Auto-map High confidence. Medium is allowed only when score is strong.
        if suggested and (confidence == "High" or (confidence == "Medium" and score >= 0.76)):
            candidates.setdefault(suggested, []).append((score, col, confidence))

    mapping = {col: "Ignore" for col in df.columns}
    for target, items in candidates.items():
        items.sort(reverse=True)
        # Only the strongest column wins for a standard field.
        mapping[items[0][1]] = target
    return mapping, suggestions


def enrich_timesheet_fields(raw_df, mapped_df, meta):
    mapped = mapped_df.copy()
    derived = []
    weekdays = meta.get("weekdays", {})
    std_col = meta.get("std_col")
    rate_col = meta.get("rate_col")
    factor_col = meta.get("factor_col")

    total_hours = None
    if weekdays:
        day_values = []
        for col in weekdays.values():
            day_values.append(pd.to_numeric(raw_df[col], errors="coerce").fillna(0.0))
        if day_values:
            total_hours = sum(day_values)
            mapped["Total Hours"] = total_hours.round(2)
            derived.append("Total Hours")

    standard_hours = None
    if std_col:
        standard_hours = pd.to_numeric(raw_df[std_col], errors="coerce").fillna(0.0)
        mapped["Standard Hours"] = standard_hours.round(2)
        derived.append("Standard Hours")

    if "Overtime Hours" not in mapped.columns and total_hours is not None and standard_hours is not None:
        mapped["Overtime Hours"] = (total_hours - standard_hours).clip(lower=0).round(2)
        derived.append("Overtime Hours")

    if rate_col:
        rate = pd.to_numeric(raw_df[rate_col], errors="coerce").fillna(0.0)
        mapped["Hourly Rate"] = rate.round(2)
        derived.append("Hourly Rate")
    else:
        rate = None

    if factor_col:
        factor = pd.to_numeric(raw_df[factor_col], errors="coerce").fillna(1.0)
        mapped["Overtime Multiplier"] = factor.round(3)
        derived.append("Overtime Multiplier")
    else:
        factor = None

    if "Overtime Pay" not in mapped.columns and "Overtime Hours" in mapped.columns and rate is not None:
        use_factor = factor if factor is not None else 1.0
        mapped["Overtime Pay"] = (pd.to_numeric(mapped["Overtime Hours"], errors="coerce").fillna(0.0) * rate * use_factor).round(2)
        derived.append("Overtime Pay")

    return clean_columns(mapped), derived


def auto_detect_and_prepare(df, file_name, key_prefix, role="input"):
    source_type, source_meta = detect_source_type(df)
    base_mapping, suggestions = build_auto_mapping(df)

    # Allow advanced overrides but keep them out of the normal workflow.
    mapping = dict(base_mapping)
    options = ["Ignore"] + CARRR_MAPPING_FIELDS

    # Initial mapped result to determine whether essential identification is understood.
    initial = apply_mapping(df, mapping)
    if source_type == "Timesheet":
        initial, initial_derived = enrich_timesheet_fields(df, initial, source_meta)
    else:
        initial_derived = []

    has_id = "Employee ID" in initial.columns
    has_name = "Employee Name" in initial.columns
    payroll_values = [c for c in RECON_FIELDS + ["Overtime Hours"] if c in initial.columns]

    uncertain = [
        (col, suggested, confidence, score)
        for col, suggested, confidence, score in suggestions
        if confidence != "High" and suggested is not None
    ]
    unmapped = [col for col, suggested, confidence, score in suggestions if mapping.get(col) == "Ignore"]

    if has_id and has_name and (payroll_values or source_type == "Timesheet"):
        overall_confidence = "High"
    elif has_id and has_name:
        overall_confidence = "Medium"
    else:
        overall_confidence = "Needs confirmation"

    st.markdown("#### CARRR Auto Detect")
    a1, a2, a3, a4 = st.columns(4)
    a1.metric("Source detected", source_type)
    a2.metric("Rows found", len(df))
    a3.metric("Fields understood", len([v for v in mapping.values() if v != "Ignore"]) + len(initial_derived))
    a4.metric("Detection", overall_confidence)

    detected_names = [v for v in mapping.values() if v != "Ignore"] + initial_derived
    if detected_names:
        st.success("CARRR understood: " + ", ".join(dict.fromkeys(detected_names)))

    # Timesheet components are calculation inputs, not payroll fields the user should map manually.
    timesheet_component_cols = set()
    if source_type == "Timesheet":
        timesheet_component_cols.update(source_meta.get("weekdays", {}).values())
        for _c in [source_meta.get("std_col"), source_meta.get("rate_col"), source_meta.get("factor_col")]:
            if _c:
                timesheet_component_cols.add(_c)

    need_advanced = not (has_id and has_name)
    with st.expander("Advanced · Review detected fields", expanded=need_advanced):
        st.caption("Normally you can leave this closed. CARRR uses timesheet calculation columns automatically and only exposes payroll fields that may need an override.")
        visible_rows = 0
        for i, (client_col, suggested, confidence, score) in enumerate(suggestions):
            if client_col in timesheet_component_cols:
                continue
            visible_rows += 1
            left, middle, right = st.columns([1.7, 1.5, .65])
            left.markdown(f"**{client_col}**")
            default_value = mapping.get(client_col, "Ignore")
            selected = middle.selectbox(
                "Detected field",
                options,
                index=options.index(default_value) if default_value in options else 0,
                key=f"v25_{key_prefix}_auto_{i}",
                label_visibility="collapsed",
            )
            right.markdown(f"{confidence}")
            mapping[client_col] = selected
        if source_type == "Timesheet" and timesheet_component_cols:
            st.caption("Automatically used for timesheet calculations: " + ", ".join(sorted(timesheet_component_cols)))
        if visible_rows == 0:
            st.caption("No payroll-field overrides are needed for this source.")

    # Resolve duplicate targets automatically. Prefer the column chosen by CARRR's
    # base auto-detection; otherwise keep the highest-confidence candidate.
    chosen = [v for v in mapping.values() if v != "Ignore"]
    duplicates = sorted({v for v in chosen if chosen.count(v) > 1 and v not in ["First Name", "Last Name"]})
    if duplicates:
        score_lookup = {col: score for col, suggested, confidence, score in suggestions}
        resolved = []
        for target in duplicates:
            cols = [col for col, value in mapping.items() if value == target]
            preferred = [col for col in cols if base_mapping.get(col) == target]
            winner = preferred[0] if preferred else max(cols, key=lambda c: score_lookup.get(c, 0.0))
            for col in cols:
                if col != winner:
                    mapping[col] = "Ignore"
            resolved.append(f"{target} → {winner}")
        st.caption("CARRR automatically resolved duplicate field matches: " + "; ".join(resolved))

    mapped = apply_mapping(df, mapping)
    derived = []
    if source_type == "Timesheet":
        mapped, derived = enrich_timesheet_fields(df, mapped, source_meta)

    missing = []
    if "Employee ID" not in mapped.columns:
        missing.append("Employee ID")
    if "Employee Name" not in mapped.columns:
        missing.append("Employee Name")

    if missing:
        st.warning("CARRR needs one quick confirmation before continuing: " + ", ".join(missing) + ". Open Advanced · Review detected fields.")
        return None, {"source_type": source_type, "confidence": "Needs confirmation", "mapping": mapping}

    if role == "output":
        available_values = [f for f in RECON_FIELDS if f in mapped.columns]
        if not available_values:
            st.warning("CARRR identified the employee columns, but could not confidently identify a payroll value such as Gross Pay, PAYE, Pension or Net Pay. Review the detected fields under Advanced.")
            return None, {"source_type": source_type, "confidence": "Needs confirmation", "mapping": mapping}

    if derived:
        st.caption("Automatically calculated from the source: " + ", ".join(derived))

    return mapped, {
        "source_type": source_type,
        "confidence": overall_confidence,
        "mapping": mapping,
        "derived": derived,
        "unmapped": unmapped,
    }

def render_mapping_editor(df, file_name, key_prefix):
    suggestions = []
    low_or_medium = False
    for col in df.columns:
        suggested, confidence, score = suggest_carrr_field(col)
        if confidence != "High":
            low_or_medium = True
        suggestions.append((col, suggested, confidence, score))

    with st.expander(f"Smart Mapping — {file_name}", expanded=low_or_medium):
        st.caption("CARRR has suggested a standard payroll field for each client column. Confirm any uncertain mappings before continuing.")
        mapping = {}
        rows = []
        options = ["Ignore"] + CARRR_MAPPING_FIELDS
        for i, (client_col, suggested, confidence, score) in enumerate(suggestions):
            left, middle, right = st.columns([1.6, 1.6, .7])
            left.markdown(f"**{client_col}**")
            default_value = suggested if suggested in options else "Ignore"
            selected = middle.selectbox(
                "CARRR field",
                options,
                index=options.index(default_value),
                key=f"{key_prefix}_map_{i}",
                label_visibility="collapsed",
            )
            right.markdown(f"**{confidence}**")
            mapping[client_col] = selected
            rows.append({"Client Header": client_col, "CARRR Field": selected, "Confidence": confidence, "Score": score})

        chosen = [v for v in mapping.values() if v != "Ignore"]
        duplicates = sorted({v for v in chosen if chosen.count(v) > 1 and v not in ["First Name", "Last Name"]})
        mapped_targets = set(chosen)
        has_name = "Employee Name" in mapped_targets or ({"First Name", "Last Name"} <= mapped_targets)
        missing_keys = []
        if "Employee ID" not in mapped_targets:
            missing_keys.append("Employee ID")
        if not has_name:
            missing_keys.append("Employee Name (or First Name + Last Name)")

        if duplicates:
            st.error("Map only one client column to each CARRR field: " + ", ".join(duplicates))
        if missing_keys:
            st.warning("Required mapping still needed: " + ", ".join(missing_keys))

        all_high = all(r[2] == "High" for r in suggestions if r[1] is not None)
        approve_default = (not low_or_medium) and not duplicates and not missing_keys
        approved = st.checkbox(
            "Approve this mapping",
            value=approve_default,
            key=f"{key_prefix}_approved",
            help="High-confidence mappings can be pre-approved. Review anything marked Medium or Low before approval.",
        )

    if duplicates or missing_keys or not approved:
        return None, pd.DataFrame(rows)
    return apply_mapping(df, mapping), pd.DataFrame(rows)


def combine_mapped_sources(mapped_sources):
    """Combine a main payroll file with supplemental structured sources using Employee ID where possible."""
    if not mapped_sources:
        return pd.DataFrame()
    frames = [df.copy() for _, df in mapped_sources]
    # Start with the richest table as the base.
    frames.sort(key=lambda d: (len(d.columns), len(d)), reverse=True)
    base = frames[0].copy()

    for extra in frames[1:]:
        if "Employee ID" in base.columns and "Employee ID" in extra.columns:
            base_ids = set(safe_text(base["Employee ID"])) - {""}
            extra_ids = set(safe_text(extra["Employee ID"])) - {""}
            overlap = len(base_ids & extra_ids) / max(1, min(len(base_ids), len(extra_ids)))
            if overlap >= 0.25:
                b = base.copy()
                e = extra.copy()
                b["Employee ID"] = safe_text(b["Employee ID"])
                e["Employee ID"] = safe_text(e["Employee ID"])
                # Avoid one-to-many joins from duplicate supplemental IDs.
                e = e.drop_duplicates(subset=["Employee ID"], keep="last")
                merged = b.merge(e, on="Employee ID", how="outer", suffixes=("", "__extra"))
                for col in list(merged.columns):
                    if col.endswith("__extra"):
                        original = col[:-7]
                        if original in merged.columns:
                            left = merged[original]
                            right = merged[col]
                            blank = left.isna() | left.astype(str).str.strip().isin(["", "nan", "None"])
                            merged.loc[blank, original] = right.loc[blank]
                            merged = merged.drop(columns=[col])
                        else:
                            merged = merged.rename(columns={col: original})
                base = merged
                continue
        # If files represent different employee populations, append rows using the union of columns.
        base = pd.concat([base, extra], ignore_index=True, sort=False)

    return clean_columns(base.reset_index(drop=True))

def _clean_pdf_header(value):
    """Repair a few common text-table extraction artefacts without changing normal headers."""
    value = str(value or "").replace("\n", " ").strip()
    value = re.sub(r"\s+", " ", value)
    low = normalize_header(value)
    # Common split-word artefacts from narrow PDF table cells.
    if low.startswith("pension contributi"):
        return "Pension Contribution"
    if low.startswith("onother deduction") or low.startswith("other deduction"):
        return "Other Deduction"
    return value


def _pdf_table_to_dataframe(uploaded_file):
    """Extract and combine compatible text-based payroll tables across PDF pages.

    Multi-page payroll reports commonly repeat the same table header on each page.
    Earlier versions selected only the single "best" table, which could make the
    last employees look missing. v2.10 combines all compatible page tables first.
    """
    raw_bytes = uploaded_file.getvalue()
    candidates = []
    pdf_pages_total = 0

    if pdfplumber is not None:
        try:
            with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
                pdf_pages_total = len(pdf.pages)

                for page_no, page in enumerate(pdf.pages[:50], start=1):
                    for table_no, table in enumerate(page.extract_tables() or [], start=1):
                        if not table or len(table) < 2:
                            continue

                        width = max(len(r) for r in table if r)
                        rows = []
                        for r in table:
                            vals = list(r or [])
                            vals += [""] * (width - len(vals))
                            rows.append([str(v or "").strip() for v in vals])

                        raw_df = pd.DataFrame(rows)
                        try:
                            df, detected = dataframe_from_raw(raw_df)
                        except Exception:
                            continue

                        df.columns = [_clean_pdf_header(c) for c in df.columns]
                        df = clean_columns(df)

                        if not len(df) or len(df.columns) < 2:
                            continue

                        recognised = 0
                        mapped_targets = []
                        for c in df.columns:
                            suggested, confidence, score = suggest_carrr_field(c)
                            if suggested and confidence in ["High", "Medium"]:
                                recognised += 1
                                mapped_targets.append(suggested)
                            else:
                                mapped_targets.append("")

                        header_signature = tuple(normalize_header(c) for c in df.columns)
                        target_signature = tuple(mapped_targets)

                        candidates.append({
                            "recognised": recognised,
                            "rows": len(df),
                            "page": page_no,
                            "table": table_no,
                            "df": df,
                            "detected_header_row": detected,
                            "header_signature": header_signature,
                            "target_signature": target_signature,
                        })
        except Exception:
            candidates = []

    if candidates:
        # Choose the strongest table schema, then combine all page tables that
        # clearly share that schema. Exact headers are preferred; mapped CARRR
        # targets provide a fallback for small PDF extraction variations.
        candidates.sort(
            key=lambda x: (x["recognised"], len(x["df"].columns), x["rows"]),
            reverse=True,
        )
        anchor = candidates[0]

        compatible = []
        for c in candidates:
            exact_schema = c["header_signature"] == anchor["header_signature"]
            mapped_schema = (
                len(c["target_signature"]) == len(anchor["target_signature"])
                and c["target_signature"] == anchor["target_signature"]
                and sum(bool(x) for x in c["target_signature"]) >= 3
            )
            if exact_schema or mapped_schema:
                page_df = c["df"].copy()

                # If the compatible schema matched through mapped targets but
                # the literal PDF headers differ slightly, align to the anchor.
                if len(page_df.columns) == len(anchor["df"].columns):
                    page_df.columns = list(anchor["df"].columns)

                compatible.append((c["page"], c["table"], page_df))

        if compatible:
            compatible.sort(key=lambda x: (x[0], x[1]))
            combined = pd.concat([x[2] for x in compatible], ignore_index=True)

            # Remove completely blank rows and exact duplicates introduced by
            # unusual PDF exporters, while preserving legitimate employees.
            combined = combined.replace(r"^\s*$", pd.NA, regex=True)
            combined = combined.dropna(how="all").fillna("")
            combined = combined.drop_duplicates().reset_index(drop=True)

            combined.attrs["detected_header_row"] = anchor["detected_header_row"]
            combined.attrs["pdf_pages_total"] = pdf_pages_total
            combined.attrs["pdf_tables_combined"] = len(compatible)
            combined.attrs["pdf_pages_with_tables"] = len(set(x[0] for x in compatible))
            combined.attrs["pdf_extracted_rows"] = len(combined)
            combined.attrs["pdf_recognised_fields"] = anchor["recognised"]
            combined.attrs["source_extraction_method"] = "Multi-page PDF table extraction"
            return combined

    # Fallback for simple PDFs where text is emitted as:
    # header1, header2, ..., value1, value2, ...
    if PdfReader is not None:
        try:
            reader = PdfReader(io.BytesIO(raw_bytes))
            lines = []
            for page in reader.pages[:20]:
                lines.extend([
                    re.sub(r"\s+", " ", line).strip()
                    for line in (page.extract_text() or "").splitlines()
                    if line.strip()
                ])

            best = None
            for i in range(min(len(lines), 120)):
                header_fields = []
                seen_targets = set()

                for j in range(i, min(i + 24, len(lines))):
                    suggested, confidence, score = suggest_carrr_field(lines[j])
                    if suggested and confidence in ["High", "Medium"] and suggested not in seen_targets:
                        header_fields.append((j, lines[j], suggested))
                        seen_targets.add(suggested)
                    else:
                        if len(header_fields) >= 3:
                            break
                        header_fields = []
                        seen_targets = set()

                if len(header_fields) >= 3:
                    h_start = header_fields[0][0]
                    h_end = header_fields[-1][0]

                    if [x[0] for x in header_fields] == list(range(h_start, h_end + 1)):
                        ncols = len(header_fields)
                        data_start = h_end + 1
                        remaining = lines[data_start:]
                        nrows = len(remaining) // ncols

                        if nrows >= 2:
                            usable = remaining[:nrows * ncols]
                            rows = [
                                usable[k:k+ncols]
                                for k in range(0, len(usable), ncols)
                            ]
                            headers = [_clean_pdf_header(x[1]) for x in header_fields]
                            df = pd.DataFrame(rows, columns=headers)

                            first_suggested = header_fields[0][2]
                            if first_suggested == "Employee ID":
                                mask = df.iloc[:, 0].astype(str).str.match(
                                    r"^[A-Za-z0-9_-]{2,}$", na=False
                                )
                                df = df.loc[mask].reset_index(drop=True)

                            if len(df) >= 2:
                                score = len(header_fields)
                                if best is None or score > best[0]:
                                    best = (score, df)

            if best is not None:
                result = clean_columns(best[1])
                result.attrs["pdf_pages_total"] = len(reader.pages)
                result.attrs["pdf_tables_combined"] = 0
                result.attrs["pdf_extracted_rows"] = len(result)
                result.attrs["source_extraction_method"] = "PDF text fallback"
                return result
        except Exception:
            pass

    raise ValueError(
        "No reliable text table was detected in this PDF. If it is scanned/image-based, "
        "CARRR will move it to OCR + human verification."
    )



def _configure_tesseract():
    """Locate the Tesseract executable when available."""
    if pytesseract is None:
        return False, "Python package pytesseract is not installed."

    try:
        pytesseract.get_tesseract_version()
        return True, ""
    except Exception:
        pass

    candidate_paths = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "Tesseract-OCR", "tesseract.exe"),
    ]
    for candidate in candidate_paths:
        if candidate and os.path.exists(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
            try:
                pytesseract.get_tesseract_version()
                return True, ""
            except Exception:
                pass

    return False, (
        "Tesseract OCR is not installed on this computer. "
        "Install Tesseract OCR, then restart CARRR. "
        "On Streamlit Cloud, packages.txt in this package installs it automatically."
    )


def _prepare_image_for_ocr(image):
    if ImageOps is None:
        return image
    img = image.convert("L")
    img = ImageOps.autocontrast(img)
    # Upscale smaller scans for better OCR.
    if img.width < 1800:
        scale = 1800 / max(img.width, 1)
        img = img.resize((int(img.width * scale), int(img.height * scale)))
    if ImageEnhance is not None:
        img = ImageEnhance.Sharpness(img).enhance(1.5)
        img = ImageEnhance.Contrast(img).enhance(1.15)
    return img


def _render_pdf_pages_for_ocr(uploaded_file, max_pages=5):
    if fitz is None or Image is None:
        raise ValueError("Scanned PDF OCR requires PyMuPDF and Pillow.")
    raw = uploaded_file.getvalue()
    pdf = fitz.open(stream=raw, filetype="pdf")
    images = []
    for page_index in range(min(len(pdf), max_pages)):
        page = pdf.load_page(page_index)
        pix = page.get_pixmap(matrix=fitz.Matrix(2.2, 2.2), alpha=False)
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        images.append(image)
    pdf.close()
    return images


def _ocr_tokens(image):
    ready, reason = _configure_tesseract()
    if not ready:
        raise ValueError(reason)

    prepared = _prepare_image_for_ocr(image)

    # PSM 11 treats the page as sparse text. It is much more reliable for
    # printed/scanned tables because grid lines do not force unrelated cells
    # into one OCR "sentence".
    attempts = []
    for psm in [11, 6]:
        try:
            data = pytesseract.image_to_data(
                prepared,
                output_type=TesseractOutput.DATAFRAME,
                config=f"--oem 3 --psm {psm}",
            )
            if data is None or data.empty:
                continue
            data = data.copy()
            data["text"] = data["text"].fillna("").astype(str).str.strip()
            data["conf"] = pd.to_numeric(data["conf"], errors="coerce").fillna(-1)
            data = data[(data["text"] != "") & (data["conf"] >= 15)].copy()
            if data.empty:
                continue

            useful = data["text"].str.contains(r"[A-Za-z0-9]", regex=True, na=False).sum()
            mean_conf = float(data.loc[
                data["text"].str.contains(r"[A-Za-z0-9]", regex=True, na=False), "conf"
            ].mean() or 0)
            # Sparse mode gets a small preference for tables.
            quality = useful + (mean_conf / 10.0) + (4 if psm == 11 else 0)
            attempts.append((quality, data))
        except Exception:
            continue

    if not attempts:
        raise ValueError("OCR detected no usable text.")

    attempts.sort(key=lambda x: x[0], reverse=True)
    return attempts[0][1]


def _line_groups_from_tokens(tokens):
    groups = []
    group_cols = ["block_num", "par_num", "line_num"]
    for _, line in tokens.groupby(group_cols, sort=False):
        line = line.sort_values("left").copy()
        words = []
        for _, row in line.iterrows():
            words.append({
                "text": str(row["text"]).strip(),
                "left": float(row["left"]),
                "top": float(row["top"]),
                "width": float(row["width"]),
                "height": float(row["height"]),
                "conf": float(row["conf"]),
            })
        if words:
            groups.append(words)
    groups.sort(key=lambda ws: (min(w["top"] for w in ws), min(w["left"] for w in ws)))
    return groups


def _scan_phrase_field(phrase):
    """Map a spatial OCR phrase to a CARRR field or a useful timesheet component."""
    suggested, confidence, score = suggest_carrr_field(phrase)
    if suggested and (confidence == "High" or score >= 0.76):
        return suggested, score

    normalized = normalize_header(phrase)
    tokens = set(normalized.split())

    weekday_aliases = {
        "Mon": {"mon", "monday"},
        "Tue": {"tue", "tues", "tuesday"},
        "Wed": {"wed", "weds", "wednesday"},
        "Thu": {"thu", "thur", "thurs", "thursday"},
        "Fri": {"fri", "friday"},
        "Sat": {"sat", "saturday"},
        "Sun": {"sun", "sunday"},
    }
    for label, aliases in weekday_aliases.items():
        if len(tokens) == 1 and tokens & aliases:
            return label, 0.97

    if "total" in tokens and tokens & {"hour", "hours", "hrs"}:
        return "Total Hours", 0.96
    if "pay" in tokens and "rate" in tokens:
        return "Pay Rate", 0.96
    if "regular" in tokens and "pay" in tokens:
        return "Regular Pay", 0.94
    if ("ot" in tokens or "overtime" in tokens) and tokens & {"hour", "hours", "hrs"}:
        return "Overtime Hours", 0.98
    if ("ot" in tokens or "overtime" in tokens) and "pay" in tokens:
        return "Overtime Pay", 0.98
    if "gross" in tokens and "pay" in tokens:
        return "Gross Pay", 0.98

    return None, 0.0


def _meaningful_ocr_tokens(tokens):
    work = tokens.copy()
    work = work[
        work["text"].astype(str).str.contains(r"[A-Za-z0-9£€$.-]", regex=True, na=False)
    ].copy()
    work["cx"] = work["left"] + (work["width"] / 2.0)
    work["cy"] = work["top"] + (work["height"] / 2.0)
    return work


def _header_candidates_in_band(tokens, y_start, y_end):
    """Recognise fields across multiple OCR lines inside one visual header band."""
    band = tokens[(tokens["top"] >= y_start) & (tokens["top"] <= y_end)].copy()
    band = band[
        band["text"].astype(str).str.contains(r"[A-Za-z0-9]", regex=True, na=False)
    ]
    if band.empty:
        return []

    items = band[["text", "left", "top", "width", "height", "conf"]].to_dict("records")
    candidates = []

    # Generate short spatial phrases. This allows:
    #   Emp + ID
    #   Employee + Name
    #   OT (top line) + Hours (bottom line)
    #   Pay + Rate
    # while rejecting combinations spread across unrelated columns.
    for size in [1, 2, 3]:
        for combo in itertools.combinations(items, size):
            left = min(x["left"] for x in combo)
            right = max(x["left"] + x["width"] for x in combo)
            top = min(x["top"] for x in combo)
            bottom = max(x["top"] + x["height"] for x in combo)
            centers = [x["left"] + x["width"] / 2.0 for x in combo]

            if right - left > 220 or bottom - top > 75:
                continue
            if size > 1 and max(centers) - min(centers) > 185:
                continue

            # Use visual reading order. A second line below a word is read after it.
            ordered = sorted(combo, key=lambda x: (round(x["top"] / 18.0), x["left"]))
            phrase = " ".join(x["text"] for x in ordered).strip()
            field, semantic_score = _scan_phrase_field(phrase)
            if not field:
                continue

            avg_conf = sum(float(x["conf"]) for x in combo) / size
            width = right - left
            candidates.append({
                "field": field,
                "text": phrase,
                "left": float(left),
                "right": float(right),
                "top": float(top),
                "bottom": float(bottom),
                "center": float((left + right) / 2.0),
                "score": float(semantic_score),
                "ocr_conf": float(avg_conf),
                "token_count": size,
                "width": float(width),
            })

    # Keep one strong, compact candidate per field. Compactness is important:
    # it prevents the OT column next to Pay Rate being mistaken for "OT Pay".
    selected = {}
    for c in candidates:
        field = c["field"]
        rank = (
            c["score"],
            c["ocr_conf"] / 100.0,
            -c["width"] / 220.0,
            -c["token_count"] / 3.0,
        )
        if field not in selected or rank > selected[field][0]:
            selected[field] = (rank, c)

    result = [item[1] for item in selected.values()]
    result.sort(key=lambda x: x["center"])
    return result


def _infer_missing_weekday_columns(fields):
    """Infer a missing weekday position when neighbouring weekday columns are clear."""
    order = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    by_day = {f["field"]: f for f in fields if f["field"] in order}
    if len(by_day) < 3:
        return fields

    unit_spacings = []
    known = [(order.index(day), item["center"]) for day, item in by_day.items()]
    known.sort()
    for (i1, c1), (i2, c2) in zip(known[:-1], known[1:]):
        if i2 > i1:
            unit = (c2 - c1) / (i2 - i1)
            if 35 <= unit <= 160:
                unit_spacings.append(unit)

    if not unit_spacings:
        return fields

    spacing = float(pd.Series(unit_spacings).median())
    anchor_index, anchor_center = known[0]
    inferred = list(fields)

    for i, day in enumerate(order):
        if day in by_day:
            continue
        predicted = anchor_center + ((i - anchor_index) * spacing)

        # Do not create a weekday column on top of an unrelated detected column.
        if any(abs(predicted - f["center"]) < spacing * 0.38 for f in fields):
            continue

        inferred.append({
            "field": day,
            "text": day,
            "left": predicted - 20,
            "right": predicted + 20,
            "top": 0,
            "bottom": 0,
            "center": predicted,
            "score": 0.80,
            "ocr_conf": 0.0,
            "token_count": 0,
            "width": 40,
            "inferred": True,
        })

    inferred.sort(key=lambda x: x["center"])
    return inferred


def _find_spatial_header(tokens):
    work = _meaningful_ocr_tokens(tokens)
    if work.empty:
        return None, None

    # Scan overlapping visual bands. Paper timesheet headers often occupy 2-3
    # printed lines, so requiring one OCR line is too strict.
    min_y = int(max(0, work["top"].quantile(0.02)))
    max_y = int(work["top"].quantile(0.55))
    best = None

    for y_start in range(min_y, max_y + 1, 10):
        fields = _header_candidates_in_band(work, y_start, y_start + 80)
        names = {f["field"] for f in fields}

        has_id = "Employee ID" in names
        has_name = "Employee Name" in names
        identification = has_id and has_name

        standard_payroll = len(
            names & {
                "Overtime Hours", "Overtime Pay", "Bonus", "Gross Pay",
                "Net Pay", "PAYE", "Employee NI", "Pension", "Salary"
            }
        )
        weekdays = len(names & {"Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"})
        useful_components = len(names & {"Pay Rate", "Regular Pay", "Total Hours"})

        score = (
            len(names)
            + (8 if identification else 0)
            + (standard_payroll * 1.8)
            + (weekdays * 0.7)
            + (useful_components * 0.5)
        )

        if identification and (standard_payroll >= 1 or weekdays >= 3):
            if best is None or score > best[0]:
                best = (score, y_start, fields)

    if best is None:
        return None, None

    fields = _infer_missing_weekday_columns(best[2])

    # The bottom of the real detected header, not inferred columns.
    detected_bottoms = [
        f["bottom"] for f in fields
        if not f.get("inferred") and f.get("bottom", 0) > 0
    ]
    header_bottom = max(detected_bottoms) if detected_bottoms else best[1] + 80
    return fields, header_bottom


def _cluster_data_rows(tokens, header_bottom):
    """Cluster sparse OCR tokens into visual employee rows by vertical position."""
    work = _meaningful_ocr_tokens(tokens)
    work = work[work["top"] > header_bottom + 8].copy()
    if work.empty:
        return []

    work = work.sort_values(["cy", "left"])
    clusters = []
    current = []
    current_mean = None

    # Employee rows in a scanned table normally have a much larger gap between
    # rows than the small top/bottom OCR variation inside one row.
    tolerance = max(14.0, min(22.0, float(work["height"].median() or 18) * 0.9))

    for rec in work.to_dict("records"):
        cy = float(rec["cy"])
        if not current:
            current = [rec]
            current_mean = cy
            continue

        if abs(cy - current_mean) <= tolerance:
            current.append(rec)
            current_mean = sum(float(x["cy"]) for x in current) / len(current)
        else:
            clusters.append(current)
            current = [rec]
            current_mean = cy

    if current:
        clusters.append(current)
    return clusters


def _table_from_ocr_tokens(tokens):
    """Reconstruct a payroll/timesheet table from OCR positions, including multi-line headers."""
    fields, header_bottom = _find_spatial_header(tokens)
    if not fields:
        raise ValueError(
            "OCR text was found, but CARRR could not confidently identify a payroll/timesheet header row."
        )

    fields = sorted(fields, key=lambda x: x["center"])
    centers = [float(f["center"]) for f in fields]
    headers = [f["field"] for f in fields]

    # Midpoints between recognised/inferred columns form robust cell boundaries.
    bounds = [-float("inf")]
    for a, b in zip(centers[:-1], centers[1:]):
        bounds.append((a + b) / 2.0)
    bounds.append(float("inf"))

    row_clusters = _cluster_data_rows(tokens, header_bottom)
    rows = []
    row_confidences = []

    id_index = headers.index("Employee ID") if "Employee ID" in headers else None
    name_index = headers.index("Employee Name") if "Employee Name" in headers else None

    for cluster in row_clusters:
        cells = [[] for _ in headers]
        cell_confs = [[] for _ in headers]

        for word in cluster:
            center = float(word["left"] + word["width"] / 2.0)
            col_idx = None
            for j in range(len(headers)):
                if bounds[j] <= center < bounds[j + 1]:
                    col_idx = j
                    break
            if col_idx is not None:
                cells[col_idx].append(word)
                cell_confs[col_idx].append(float(word["conf"]))

        values = []
        for group in cells:
            group = sorted(group, key=lambda x: x["left"])
            values.append(" ".join(str(x["text"]).strip() for x in group).strip())

        nonempty = sum(bool(v) for v in values)
        if nonempty < 2:
            continue

        employee_id = values[id_index] if id_index is not None else ""
        employee_name = values[name_index] if name_index is not None else ""

        # Filter signatures/footer text while allowing imperfect OCR IDs such as EMPOO1.
        compact_id = re.sub(r"[^A-Za-z0-9_-]", "", employee_id)
        id_like = bool(
            compact_id
            and re.search(r"[A-Za-z]", compact_id)
            and re.search(r"\d", compact_id)
            and len(compact_id) <= 24
        )
        name_like = bool(re.search(r"[A-Za-z]{2}", employee_name or ""))

        payroll_values = sum(
            bool(values[headers.index(h)])
            for h in ["Overtime Hours", "Overtime Pay", "Bonus", "Gross Pay", "Net Pay"]
            if h in headers
        )

        if not id_like and not (name_like and payroll_values >= 2):
            continue

        rows.append(values)
        conf_values = [c for group in cell_confs for c in group]
        row_confidences.append(
            sum(conf_values) / len(conf_values) if conf_values else 0.0
        )

    if not rows:
        raise ValueError(
            "CARRR identified the scanned table header, but could not safely reconstruct employee rows."
        )

    df = pd.DataFrame(rows, columns=headers)

    # Remove columns that are entirely empty and make duplicate headers safe.
    df = df.replace("", pd.NA)
    empty_cols = [c for c in df.columns if df[c].isna().all()]
    if empty_cols:
        df = df.drop(columns=empty_cols)
    df = df.fillna("").reset_index(drop=True)

    token_conf = float(pd.to_numeric(tokens["conf"], errors="coerce").replace(-1, pd.NA).dropna().mean() or 0)
    row_conf = float(sum(row_confidences) / len(row_confidences)) if row_confidences else token_conf

    # Confidence also reflects how complete the reconstructed table is.
    nonblank_ratio = (
        float((df.astype(str).apply(lambda s: s.str.strip() != "")).sum().sum())
        / max(float(df.shape[0] * df.shape[1]), 1.0)
    )
    adjusted_conf = max(0.0, min(100.0, (row_conf * 0.72) + (nonblank_ratio * 28.0)))

    df.attrs["ocr_confidence"] = adjusted_conf
    df.attrs["ocr_header_fields"] = list(df.columns)
    df.attrs["ocr_spatial_header"] = True
    return clean_columns(df)


def extract_scan_to_dataframe(uploaded_file):
    """OCR a photo or image-only PDF into a candidate structured table."""
    name = uploaded_file.name.lower()

    if name.endswith((".png", ".jpg", ".jpeg")):
        if Image is None:
            raise ValueError("Image OCR requires Pillow.")
        image = Image.open(io.BytesIO(uploaded_file.getvalue()))
        tokens = _ocr_tokens(image)
        return _table_from_ocr_tokens(tokens)

    if name.endswith(".pdf"):
        images = _render_pdf_pages_for_ocr(uploaded_file)
        page_tables = []
        for image in images:
            try:
                tokens = _ocr_tokens(image)
                page_tables.append(_table_from_ocr_tokens(tokens))
            except Exception:
                continue
        if not page_tables:
            raise ValueError("OCR could not reconstruct a reliable table from this scanned PDF.")
        # Keep identical-column page tables together.
        base_cols = list(page_tables[0].columns)
        compatible = [df for df in page_tables if list(df.columns) == base_cols]
        result = pd.concat(compatible, ignore_index=True)
        confidences = [df.attrs.get("ocr_confidence", 0) for df in compatible]
        result.attrs["ocr_confidence"] = sum(confidences) / len(confidences) if confidences else 0
        return clean_columns(result)

    raise ValueError("OCR is only used for images and scanned PDFs.")


def render_ocr_verification(uploaded_file, key_prefix):
    """Human-in-the-loop verification before OCR data is permitted into CARRR."""
    st.markdown("#### CARRR Document Verification · Spatial OCR")
    st.caption(
        "CARRR extracted a candidate table from this scan/photo. "
        "Review the values below. OCR data is not used until you explicitly approve it."
    )

    cache_key = f"{key_prefix}_ocr_candidate"
    if cache_key not in st.session_state:
        try:
            st.session_state[cache_key] = extract_scan_to_dataframe(uploaded_file)
        except Exception as exc:
            st.warning("CARRR could not safely structure this scan/photo automatically.")
            st.write(str(exc))
            ready, reason = _configure_tesseract()
            if not ready:
                st.info(reason)
            return None

    candidate = st.session_state[cache_key]
    confidence = candidate.attrs.get("ocr_confidence", 0.0)
    v1, v2, v3 = st.columns(3)
    v1.metric("Rows extracted", len(candidate))
    v2.metric("Columns extracted", len(candidate.columns))
    v3.metric("OCR confidence", f"{confidence:.0f}%")

    if confidence < 70:
        st.warning("OCR confidence is low. Check every value carefully before approval.")
    elif confidence < 85:
        st.info("OCR confidence is moderate. Review highlighted/uncertain values carefully.")
    else:
        st.success("OCR confidence is high, but human approval is still required.")

    edited = st.data_editor(
        candidate,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        key=f"{key_prefix}_ocr_editor",
    )

    confirm = st.checkbox(
        "I have checked the extracted values and approve this data for CARRR processing.",
        key=f"{key_prefix}_ocr_confirm",
    )

    if confirm:
        verified = clean_columns(pd.DataFrame(edited))
        verified.attrs["ocr_verified"] = True
        verified.attrs["ocr_confidence"] = confidence
        st.success("Extracted data approved. CARRR can now use it.")
        return verified

    st.info("Review the table, correct anything necessary, then tick the approval box to continue.")
    return None


def read_upload(uploaded_file):
    """Universal structured-source reader used before AND after payroll."""
    name = uploaded_file.name.lower()

    if name.endswith(".csv"):
        raw = pd.read_csv(uploaded_file, header=None, dtype=object)
        df, header_row = dataframe_from_raw(raw)
        df.attrs["detected_header_row"] = header_row
        return df

    if name.endswith(".tsv"):
        raw = pd.read_csv(uploaded_file, sep="\t", header=None, dtype=object)
        df, header_row = dataframe_from_raw(raw)
        df.attrs["detected_header_row"] = header_row
        return df

    if name.endswith(".txt"):
        # Auto-detect comma/tab/semicolon/pipe delimited text.
        raw = pd.read_csv(uploaded_file, sep=None, engine="python", header=None, dtype=object)
        df, header_row = dataframe_from_raw(raw)
        df.attrs["detected_header_row"] = header_row
        return df

    if name.endswith((".xlsx", ".xls")):
        raw = pd.read_excel(uploaded_file, header=None, dtype=object)
        df, header_row = dataframe_from_raw(raw)
        df.attrs["detected_header_row"] = header_row
        return df

    if name.endswith(".docx"):
        if Document is None:
            raise ValueError("Word support is not installed. Install python-docx.")
        doc = Document(io.BytesIO(uploaded_file.getvalue()))
        if not doc.tables:
            raise ValueError("No usable table was found in this Word document.")
        table = max(doc.tables, key=lambda t: max(len(t.rows), 1) * max(len(t.columns), 1))
        values = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(values) < 2:
            raise ValueError("The Word table does not contain enough rows.")
        raw = pd.DataFrame(values)
        df, header_row = dataframe_from_raw(raw)
        df.attrs["detected_header_row"] = header_row
        return df

    if name.endswith(".pdf"):
        return _pdf_table_to_dataframe(uploaded_file)

    if name.endswith(".json"):
        try:
            obj = pd.read_json(io.BytesIO(uploaded_file.getvalue()))
            return clean_columns(obj)
        except Exception:
            import json
            payload = json.loads(uploaded_file.getvalue().decode("utf-8"))
            if isinstance(payload, dict):
                # Common API/export pattern: first list-valued key contains records.
                for value in payload.values():
                    if isinstance(value, list):
                        return clean_columns(pd.json_normalize(value))
                return clean_columns(pd.json_normalize(payload))
            return clean_columns(pd.json_normalize(payload))

    if name.endswith(".xml"):
        try:
            return clean_columns(pd.read_xml(io.BytesIO(uploaded_file.getvalue())))
        except Exception as exc:
            raise ValueError(f"Could not read XML table: {exc}")

    if name.endswith((".png", ".jpg", ".jpeg")):
        raise ValueError("OCR_VERIFICATION_REQUIRED")

    raise ValueError("Unsupported source type.")


def inspect_unstructured_source(uploaded_file):
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        try:
            df = _pdf_table_to_dataframe(uploaded_file)
            return {
                "type": "PDF table",
                "status": "Structured",
                "detail": f"CARRR extracted {len(df)} row(s) and {len(df.columns)} column(s).",
                "preview": "",
            }
        except Exception:
            preview = ""
            if PdfReader is not None:
                try:
                    reader = PdfReader(io.BytesIO(uploaded_file.getvalue()))
                    preview = "\n".join((page.extract_text() or "") for page in reader.pages[:5]).strip()[:3000]
                except Exception:
                    pass
            return {
                "type": "PDF / scan",
                "status": "Verification required",
                "detail": "No reliable text table was detected. CARRR can try OCR, but human confirmation is required.",
                "preview": preview,
            }

    if name.endswith((".png", ".jpg", ".jpeg")):
        return {
            "type": "Scan / photo",
            "status": "Verification required",
            "detail": "Image received. CARRR can run OCR, but the extracted values require human confirmation before use.",
            "preview": "",
        }

    return {
        "type": "Other",
        "status": "Verification required",
        "detail": "Source received but could not be converted into a structured payroll table.",
        "preview": "",
    }


def numeric(series):
    return pd.to_numeric(series, errors="coerce").fillna(0.0)


def safe_text(series):
    return series.fillna("").astype(str).str.strip()


def employee_name(df):
    if "Employee Name" in df.columns:
        return safe_text(df["Employee Name"])
    return pd.Series([""] * len(df), index=df.index)


def add_issue(rows, severity, check_name, detail, emp_id="", emp_name="", source="CHECK"):
    rows.append({
        "Source": source,
        "Status": severity,
        "Employee ID": str(emp_id) if emp_id is not None else "",
        "Employee Name": str(emp_name) if emp_name is not None else "",
        "Check / Issue": check_name,
        "Detail": detail,
    })


def check_payroll(df, label):
    rows = []
    names = employee_name(df)

    # Required columns
    for col in REQUIRED_KEYS:
        if col not in df.columns:
            add_issue(rows, "Fail", f"{label}: required column", f"Missing required column: {col}")

    if "Employee ID" not in df.columns:
        return pd.DataFrame(rows)

    ids = safe_text(df["Employee ID"])

    # Blank key data
    for idx in df.index[ids.eq("")]:
        add_issue(rows, "Fail", "Missing Employee ID", "Employee ID is blank.", "", names.loc[idx], "CHECK")

    if "Employee Name" in df.columns:
        blank_names = safe_text(df["Employee Name"]).eq("")
        for idx in df.index[blank_names]:
            add_issue(rows, "Fail", "Missing Employee Name", "Employee Name is blank.", ids.loc[idx], "", "CHECK")

    # Duplicate IDs
    dup = ids.ne("") & ids.duplicated(keep=False)
    for idx in df.index[dup]:
        add_issue(rows, "Fail", "Duplicate employee", "Duplicate Employee ID detected.", ids.loc[idx], names.loc[idx], "CHECK")

    # Optional payroll master-data checks: only enforce when the column exists.
    optional_fields = {
        "Tax Code": "Tax code is blank.",
        "NI Number": "NI number is blank.",
        "Pension": "Pension value/status is blank.",
        "Bank Details": "Bank details are blank.",
        "Sort Code": "Sort code is blank.",
        "Account Number": "Account number is blank.",
    }
    for col, detail in optional_fields.items():
        if col in df.columns:
            blank = safe_text(df[col]).eq("")
            for idx in df.index[blank]:
                add_issue(rows, "Warning", f"Missing {col}", detail, ids.loc[idx], names.loc[idx], "CHECK")

    # Light NI format check for demo validation (not a substitute for HMRC validation).
    if "NI Number" in df.columns:
        ni = safe_text(df["NI Number"]).str.upper().str.replace(" ", "", regex=False)
        valid = ni.eq("") | ni.str.match(r"^[A-CEGHJ-PR-TW-Z]{2}\d{6}[A-D]$")
        for idx in df.index[~valid]:
            add_issue(rows, "Warning", "NI number format", "NI number format looks unusual; verify against source data.", ids.loc[idx], names.loc[idx], "CHECK")

    # Numeric sanity checks for available pay fields.
    for col in ["Gross Pay", "Salary", "PAYE", "Employee NI", "Pension", "Other Deductions", "Net Pay", "Overtime", "Overtime Pay", "Overtime Hours", "Bonus"]:
        if col in df.columns:
            raw = pd.to_numeric(df[col], errors="coerce")
            bad = raw.isna() & safe_text(df[col]).ne("")
            for idx in df.index[bad]:
                add_issue(rows, "Fail", f"Invalid {col}", f"{col} is not numeric.", ids.loc[idx], names.loc[idx], "CHECK")

    if not rows:
        add_issue(rows, "Pass", f"{label}: validation", "No validation issues detected in the supplied columns.", source="CHECK")

    return pd.DataFrame(rows)


def analyse_payroll(expected, processed, overtime_limit, pay_change_pct, overtime_hours_limit):
    rows = []
    if "Employee ID" not in processed.columns:
        add_issue(rows, "Fail", "Analysis unavailable", "Processed payroll has no Employee ID column.", source="ANALYSE")
        return pd.DataFrame(rows)

    proc = processed.copy()
    proc["Employee ID"] = safe_text(proc["Employee ID"])
    proc_names = employee_name(proc)

    # Duplicate employee IDs
    dup = proc["Employee ID"].ne("") & proc["Employee ID"].duplicated(keep=False)
    for idx in proc.index[dup]:
        add_issue(rows, "Fail", "Duplicate employee", "Duplicate employee appears in processed payroll.", proc.loc[idx, "Employee ID"], proc_names.loc[idx], "ANALYSE")

    # Negative net pay
    if "Net Pay" in proc.columns:
        net = numeric(proc["Net Pay"])
        for idx in proc.index[net.lt(0)]:
            add_issue(rows, "Fail", "Negative net pay", f"Net pay is £{net.loc[idx]:,.2f}.", proc.loc[idx, "Employee ID"], proc_names.loc[idx], "ANALYSE")

    # High overtime amount.
    overtime_col = None
    if "Overtime" in proc.columns:
        overtime_col = "Overtime"
    elif "Overtime Pay" in proc.columns:
        overtime_col = "Overtime Pay"
    if overtime_col:
        ot = numeric(proc[overtime_col])
        for idx in proc.index[ot.gt(overtime_limit)]:
            add_issue(rows, "Warning", "Unusually high overtime", f"{overtime_col} is {ot.loc[idx]:,.2f}, above the review threshold of {overtime_limit:,.2f}.", proc.loc[idx, "Employee ID"], proc_names.loc[idx], "ANALYSE")

    # High overtime hours when hours are supplied separately.
    if "Overtime Hours" in proc.columns:
        ot_hours = numeric(proc["Overtime Hours"])
        for idx in proc.index[ot_hours.gt(overtime_hours_limit)]:
            add_issue(rows, "Warning", "Unusually high overtime hours", f"Overtime Hours are {ot_hours.loc[idx]:,.2f}, above the review threshold of {overtime_hours_limit:,.2f} hours.", proc.loc[idx, "Employee ID"], proc_names.loc[idx], "ANALYSE")

    # Large pay movements between expected and processed.
    pay_col = "Gross Pay" if "Gross Pay" in proc.columns and "Gross Pay" in expected.columns else ("Salary" if "Salary" in proc.columns and "Salary" in expected.columns else None)
    if pay_col and "Employee ID" in expected.columns:
        exp = expected[["Employee ID", pay_col]].copy()
        exp["Employee ID"] = safe_text(exp["Employee ID"])
        exp[f"{pay_col} Expected"] = numeric(exp[pay_col])
        merged = proc[["Employee ID", pay_col]].copy().merge(
            exp[["Employee ID", f"{pay_col} Expected"]], on="Employee ID", how="left"
        )
        merged[f"{pay_col} Processed"] = numeric(merged[pay_col])
        base = merged[f"{pay_col} Expected"].abs()
        merged["Change %"] = ((merged[f"{pay_col} Processed"] - merged[f"{pay_col} Expected"]).abs() / base.replace(0, pd.NA) * 100).fillna(0)
        big = merged["Change %"].gt(pay_change_pct)
        name_lookup = proc.set_index("Employee ID").get("Employee Name", pd.Series(dtype=object)) if "Employee Name" in proc.columns else pd.Series(dtype=object)
        for idx in merged.index[big]:
            emp_id = merged.loc[idx, "Employee ID"]
            emp_name = ""
            if not name_lookup.empty and emp_id in name_lookup.index:
                value = name_lookup.loc[emp_id]
                emp_name = value.iloc[0] if isinstance(value, pd.Series) else value
            add_issue(rows, "Warning", "Large pay change", f"{pay_col} changed by {merged.loc[idx, 'Change %']:.1f}% against expected payroll.", emp_id, emp_name, "ANALYSE")

    # Unexpected processed employees
    if "Employee ID" in expected.columns:
        expected_ids = set(safe_text(expected["Employee ID"]))
        unexpected = ~proc["Employee ID"].isin(expected_ids)
        for idx in proc.index[unexpected]:
            add_issue(rows, "Fail", "Unexpected payment / employee", "Employee appears in processed payroll but not in expected payroll.", proc.loc[idx, "Employee ID"], proc_names.loc[idx], "ANALYSE")

    # Leaver payment check when a status field exists.
    status_col = next((c for c in ["Employment Status", "Employee Status", "Starter/Leaver", "Status"] if c in proc.columns), None)
    if status_col and "Net Pay" in proc.columns:
        status = safe_text(proc[status_col]).str.lower()
        net = numeric(proc["Net Pay"])
        leaver_paid = status.str.contains("leaver|left|terminated|inactive", regex=True) & net.gt(0)
        for idx in proc.index[leaver_paid]:
            add_issue(rows, "Warning", "Leaver payment", f"{status_col} indicates a leaver/inactive employee but net pay is £{net.loc[idx]:,.2f}.", proc.loc[idx, "Employee ID"], proc_names.loc[idx], "ANALYSE")

    if not rows:
        add_issue(rows, "Pass", "Automated analysis", "No analytical exceptions detected using the selected thresholds.", source="ANALYSE")

    return pd.DataFrame(rows)


def _normalised_employee_ids(df):
    if "Employee ID" not in df.columns:
        return pd.Series(dtype=str)
    ids = safe_text(df["Employee ID"]).astype(str).str.strip()
    ids = ids.replace({"nan": "", "None": "", "<NA>": ""})
    return ids


def source_validation_summary(expected, processed, source_name, raw_source=None):
    """Validate employee-population integrity before reconciliation.

    This does not replace reconciliation. It makes population differences visible
    before CARRR calculates pay mismatches, so extraction problems cannot be silent.
    """
    expected_ids = _normalised_employee_ids(expected)
    processed_ids = _normalised_employee_ids(processed)

    expected_nonblank = expected_ids[expected_ids.ne("")]
    processed_nonblank = processed_ids[processed_ids.ne("")]

    expected_set = set(expected_nonblank.tolist())
    processed_set = set(processed_nonblank.tolist())

    missing_ids = sorted(expected_set - processed_set)
    unexpected_ids = sorted(processed_set - expected_set)

    expected_dup_rows = int(expected_nonblank.duplicated(keep=False).sum())
    processed_dup_rows = int(processed_nonblank.duplicated(keep=False).sum())
    blank_expected = int(expected_ids.eq("").sum())
    blank_processed = int(processed_ids.eq("").sum())

    population_aligned = (
        len(missing_ids) == 0
        and len(unexpected_ids) == 0
        and expected_dup_rows == 0
        and processed_dup_rows == 0
        and blank_expected == 0
        and blank_processed == 0
    )

    raw_attrs = getattr(raw_source, "attrs", {}) if raw_source is not None else {}

    summary = {
        "Expected Employees": len(expected),
        "Extracted Employees": len(processed),
        "Missing IDs": len(missing_ids),
        "Unexpected IDs": len(unexpected_ids),
        "Duplicate Expected ID Rows": expected_dup_rows,
        "Duplicate Output ID Rows": processed_dup_rows,
        "Blank Expected IDs": blank_expected,
        "Blank Output IDs": blank_processed,
        "Population Aligned": population_aligned,
        "Missing ID List": missing_ids,
        "Unexpected ID List": unexpected_ids,
        "PDF Pages": raw_attrs.get("pdf_pages_total"),
        "PDF Tables Combined": raw_attrs.get("pdf_tables_combined"),
        "PDF Pages With Tables": raw_attrs.get("pdf_pages_with_tables"),
        "Extraction Method": raw_attrs.get("source_extraction_method", ""),
        "Source Name": source_name,
    }
    return summary


def source_validation_table(summary):
    rows = [
        {"Check": "Expected employees", "Result": summary["Expected Employees"]},
        {"Check": "Extracted employees", "Result": summary["Extracted Employees"]},
        {"Check": "Missing employee IDs", "Result": summary["Missing IDs"]},
        {"Check": "Unexpected employee IDs", "Result": summary["Unexpected IDs"]},
        {"Check": "Duplicate output ID rows", "Result": summary["Duplicate Output ID Rows"]},
        {"Check": "Blank output IDs", "Result": summary["Blank Output IDs"]},
    ]

    if summary.get("PDF Pages") is not None:
        rows.extend([
            {"Check": "PDF pages", "Result": summary.get("PDF Pages", "")},
            {"Check": "PDF tables combined", "Result": summary.get("PDF Tables Combined", "")},
            {"Check": "PDF pages containing payroll tables", "Result": summary.get("PDF Pages With Tables", "")},
        ])

    if summary.get("Extraction Method"):
        rows.append({"Check": "Extraction method", "Result": summary["Extraction Method"]})

    return pd.DataFrame(rows)


def validate_reconciliation_columns(df, label, fields):
    errors = []
    required = REQUIRED_KEYS + fields
    missing = [c for c in required if c not in df.columns]
    if missing:
        errors.append(f"{label}: missing columns: {', '.join(missing)}")
    if "Employee ID" in df.columns:
        dup = safe_text(df["Employee ID"]).duplicated(keep=False)
        if dup.any():
            errors.append(f"{label}: {int(dup.sum())} rows contain duplicate Employee IDs.")
    return errors


def build_exception_reason(row, fields):
    reasons = []
    if row["Record Status"] == "Missing from processed payroll":
        reasons.append("Missing from processed payroll")
    elif row["Record Status"] == "Unexpected employee in processed payroll":
        reasons.append("Unexpected employee in processed payroll")
    for field in fields:
        if row.get(f"{field} Status") == "Mismatch":
            diff = row.get(f"{field} Difference", 0)
            reasons.append(f"{field} mismatch ({diff:+,.2f})")
    return "; ".join(reasons) if reasons else "No exception"


def reconcile(expected, processed, fields, tolerance):
    expected = expected.copy()
    processed = processed.copy()
    expected["Employee ID"] = safe_text(expected["Employee ID"])
    processed["Employee ID"] = safe_text(processed["Employee ID"])

    expected = expected[["Employee ID", "Employee Name"] + fields]
    processed = processed[["Employee ID", "Employee Name"] + fields]

    expected = expected.rename(columns={c: f"{c} Expected" for c in expected.columns if c != "Employee ID"})
    processed = processed.rename(columns={c: f"{c} Processed" for c in processed.columns if c != "Employee ID"})

    merged = expected.merge(processed, on="Employee ID", how="outer", indicator=True)
    merged["Employee Name"] = merged["Employee Name Expected"].fillna(merged["Employee Name Processed"])

    status_cols = []
    for field in fields:
        exp_col = f"{field} Expected"
        proc_col = f"{field} Processed"
        diff_col = f"{field} Difference"
        status_col = f"{field} Status"
        merged[exp_col] = numeric(merged[exp_col])
        merged[proc_col] = numeric(merged[proc_col])
        merged[diff_col] = (merged[proc_col] - merged[exp_col]).round(2)
        merged[status_col] = merged[diff_col].abs().le(tolerance).map({True: "OK", False: "Mismatch"})
        status_cols.append(status_col)

    merged["Record Status"] = merged["_merge"].map({
        "left_only": "Missing from processed payroll",
        "right_only": "Unexpected employee in processed payroll",
        "both": "Present in both",
    })
    merged["Overall Status"] = "OK"
    merged.loc[merged["_merge"] != "both", "Overall Status"] = "Exception"
    merged.loc[merged[status_cols].eq("Mismatch").any(axis=1), "Overall Status"] = "Exception"
    merged["Exception Reason"] = merged.apply(lambda r: build_exception_reason(r, fields), axis=1)

    ordered = [
        "Employee ID", "Employee Name", "Record Status", "Overall Status", "Exception Reason"
    ] + [
        col for field in fields for col in (
            f"{field} Expected", f"{field} Processed", f"{field} Difference", f"{field} Status"
        )
    ]
    return merged[ordered]


def control_totals(result, fields):
    rows = []
    for field in fields:
        exp = result[f"{field} Expected"].sum()
        proc = result[f"{field} Processed"].sum()
        rows.append({
            "Payroll Control": field,
            "Expected Total": round(exp, 2),
            "Processed Total": round(proc, 2),
            "Variance": round(proc - exp, 2),
        })
    return pd.DataFrame(rows)


def reconciliation_review_items(result):
    rows = []
    exceptions = result[result["Overall Status"].eq("Exception")]
    for _, r in exceptions.iterrows():
        rows.append({
            "Source": "RECONCILE",
            "Status": "Fail",
            "Employee ID": r.get("Employee ID", ""),
            "Employee Name": r.get("Employee Name", ""),
            "Check / Issue": "Reconciliation exception",
            "Detail": r.get("Exception Reason", ""),
        })
    if not rows:
        add_issue(rows, "Pass", "Reconciliation", "Expected and processed payroll match within tolerance.", source="RECONCILE")
    return pd.DataFrame(rows)



def control_coverage(expected, processed):
    controls = [
        ("Employee ID", ["Employee ID"]),
        ("Employee Name", ["Employee Name"]),
        ("Tax Code", ["Tax Code"]),
        ("NI Number", ["NI Number"]),
        ("Bank Details", ["Bank Details", "Sort Code", "Account Number"]),
        ("Pension", ["Pension"]),
        ("Starter / Leaver Status", ["Employment Status", "Employee Status", "Starter/Leaver", "Status"]),
        ("Salary / Gross Pay", ["Salary", "Gross Pay"]),
        ("Overtime", ["Overtime", "Overtime Pay", "Overtime Hours"]),
        ("Bonus", ["Bonus"]),
        ("Deductions", ["PAYE", "Employee NI", "Other Deductions"]),
        ("Net Pay", ["Net Pay"]),
    ]
    rows = []
    for label, candidates in controls:
        exp_cols = [c for c in candidates if c in expected.columns]
        proc_cols = [c for c in candidates if c in processed.columns]
        rows.append({
            "Control": label,
            "Expected File": ", ".join(exp_cols) if exp_cols else "Not supplied",
            "Processed File": ", ".join(proc_cols) if proc_cols else "Not supplied",
            "Coverage": "Available" if exp_cols and proc_cols else "Partial / Missing",
        })
    return pd.DataFrame(rows)

def build_review_queue(check_df, analysis_df, recon_df):
    queue = pd.concat([check_df, analysis_df, recon_df], ignore_index=True)
    queue.insert(0, "Exception ID", [f"CARRR-{i+1:04d}" for i in range(len(queue))])
    queue["Reviewed"] = queue["Status"].eq("Pass")
    queue["Review Status"] = queue["Status"].map({"Pass": "Not required", "Warning": "Open", "Fail": "Open"})
    queue["Reviewer Note"] = ""
    return queue


def to_excel_bytes(summary_df, check_df, analysis_df, controls_df, recon_df, review_df, release_df, blockers_df=None):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        check_df.to_excel(writer, sheet_name="Check", index=False)
        analysis_df.to_excel(writer, sheet_name="Analyse", index=False)
        controls_df.to_excel(writer, sheet_name="Control Totals", index=False)
        recon_df.to_excel(writer, sheet_name="Reconciliation", index=False)
        review_df.to_excel(writer, sheet_name="Review", index=False)
        release_df.to_excel(writer, sheet_name="Release", index=False)
        if blockers_df is not None:
            blockers_df.to_excel(writer, sheet_name="Release Blockers", index=False)
    return output.getvalue()




def analyse_input_payroll(df, overtime_limit=500.0, overtime_hours_limit=20.0, bonus_limit=1000.0):
    rows=[]
    if "Employee ID" not in df.columns:
        add_issue(rows,"Fail","Analysis unavailable","Payroll input has no Employee ID column.",source="ANALYSE")
        return pd.DataFrame(rows)
    ids=safe_text(df["Employee ID"]); names=employee_name(df)
    dup=ids.ne("") & ids.duplicated(keep=False)
    for idx in df.index[dup]:
        add_issue(rows,"Fail","Duplicate employee","Employee ID appears more than once in payroll input.",ids.loc[idx],names.loc[idx],"ANALYSE")
    for col,label,limit in [("Overtime Pay","High overtime pay",overtime_limit),("Overtime","High overtime",overtime_limit),("Overtime Hours","High overtime hours",overtime_hours_limit),("Bonus","High bonus",bonus_limit)]:
        if col in df.columns:
            vals=numeric(df[col])
            for idx in df.index[vals.gt(limit)]:
                add_issue(rows,"Warning",label,f"{col} is {vals.loc[idx]:,.2f}, above the review threshold of {limit:,.2f}.",ids.loc[idx],names.loc[idx],"ANALYSE")
    if not rows:
        add_issue(rows,"Pass","Before-payroll analysis","No input exceptions detected using the selected thresholds.",source="ANALYSE")
    return pd.DataFrame(rows)


def attention_only(df):
    if df is None or df.empty or "Status" not in df.columns:
        return pd.DataFrame(columns=["Status","Employee ID","Employee Name","Check / Issue","Detail"])
    return df[df["Status"].isin(["Warning","Fail"])].copy()


# -----------------------------
# Simple CARRR control-centre workflow
# -----------------------------
with st.expander("What can the CARRR Input Hub accept?", expanded=False):
    st.markdown("""
**Automatic preprocessing supported**
- Excel (`.xlsx` / `.xls`)
- CSV / TSV / delimited text
- Word (`.docx`) tables
- text-based PDF tables
- JSON / XML exports

**Verification required**
- scanned/image-only PDFs → OCR + mandatory verification
- photos / paper timesheets → OCR + mandatory verification

CARRR never treats OCR or extracted document values as approved payroll data until a user has reviewed and approved them.
""")

payroll_period=st.text_input("Payroll period",placeholder="Example: October 2026")

st.markdown("""
<div class="before-box">
 <div class="kicker">Input Hub</div>
 <div class="bigline">1. Add the client payroll sources</div>
 <div class="muted">Upload one or more files. CARRR automatically detects the source type, header row and recognised payroll fields. Manual mapping stays hidden unless something is unclear.</div>
</div>
""", unsafe_allow_html=True)

src1,src2,src3,src4=st.columns(4)
src1.markdown("**📊 Spreadsheet**<br>Excel / CSV", unsafe_allow_html=True)
src2.markdown("**📝 Word table**<br>DOCX", unsafe_allow_html=True)
src3.markdown("**📄 PDF / data export**<br>Auto extract when text-based", unsafe_allow_html=True)
src4.markdown("**📷 Scan / photo**<br>OCR + verify", unsafe_allow_html=True)

input_files=st.file_uploader(
    "Client payroll source files",
    type=["csv","tsv","txt","xlsx","xls","docx","pdf","json","xml","png","jpg","jpeg"],
    accept_multiple_files=True,
    key="carrr_input_sources"
)

if not input_files:
    st.info("Start in the Input Hub: add the client payroll source file(s).")
    st.stop()

structured_sources=[]
verification_items=[]

for file_index, source_file in enumerate(input_files):
    try:
        source_df=clean_columns(read_upload(source_file))
        structured_sources.append((source_file.name, source_df, file_index))
    except Exception as exc:
        lower_name = source_file.name.lower()
        is_image = lower_name.endswith((".png", ".jpg", ".jpeg"))
        is_pdf = lower_name.endswith(".pdf")

        # Image sources and PDFs without a reliable text table enter OCR + verification.
        if is_image or is_pdf:
            verified_scan = render_ocr_verification(source_file, f"input_ocr_{file_index}")
            if verified_scan is not None:
                structured_sources.append((source_file.name, verified_scan, file_index))
                continue

        info=inspect_unstructured_source(source_file)
        verification_items.append({
            "File": source_file.name,
            "Source Type": info["type"],
            "Status": info["status"],
            "Reason": info["detail"] if info.get("detail") else str(exc),
        })
        if info.get("preview"):
            with st.expander(f"Extracted text preview — {source_file.name}", expanded=False):
                st.text(info["preview"])

st.markdown("#### Input Hub status")
ih1,ih2,ih3=st.columns(3)
ih1.metric("Files received",len(input_files))
ih2.metric("Structured sources",len(structured_sources))
ih3.metric("Require verification",len(verification_items))

if verification_items:
    st.warning("Some documents are captured but are not yet allowed to drive payroll checks. They remain in the verification queue.")
    st.dataframe(pd.DataFrame(verification_items),use_container_width=True,hide_index=True)

if not structured_sources:
    st.info("No approved structured payroll table is ready yet. Review and approve any OCR-extracted scan/photo data above.")
    st.stop()

st.markdown("### CARRR Auto Detect")
st.caption("CARRR identifies the source and standardises recognised payroll fields automatically. Manual mapping is hidden under Advanced and is only needed for unusual client-specific headers.")

mapped_sources=[]
for file_name, source_df, file_index in structured_sources:
    st.markdown(f"**{file_name}**")
    detected_row = source_df.attrs.get("detected_header_row", 1)
    if detected_row > 1:
        st.caption(f"Header row detected automatically: row {detected_row}")
    mapped_df, detect_info = auto_detect_and_prepare(source_df, file_name, f"v25_input_{file_index}", role="input")
    if mapped_df is not None:
        mapped_sources.append((file_name, mapped_df))

if len(mapped_sources) != len(structured_sources):
    st.info("CARRR is waiting only for the unclear fields shown above. Confirm those under Advanced to continue.")
    st.stop()

expected=combine_mapped_sources(mapped_sources)
expected=clean_columns(expected)

st.success(f"Auto Detect complete. CARRR created one standard payroll input from {len(mapped_sources)} structured source file(s).")
with st.expander("Preview CARRR standardised payroll input", expanded=False):
    st.dataframe(expected.head(50),use_container_width=True,hide_index=True)

st.markdown('<div class="before-box"><div class="kicker">Before payroll</div><div class="bigline">2. Check + Analyse the verified input</div><div class="muted">CARRR now checks the standardised payroll input and surfaces only the records that need attention.</div></div>',unsafe_allow_html=True)
with st.expander("Before-payroll thresholds", expanded=False):
    t1,t2,t3=st.columns(3)
    overtime_limit=t1.number_input("High overtime pay (£)",min_value=0.0,value=500.0,step=50.0)
    overtime_hours_limit=t2.number_input("High overtime hours",min_value=0.0,value=20.0,step=1.0)
    bonus_limit=t3.number_input("High bonus (£)",min_value=0.0,value=1000.0,step=100.0)

check_input=check_payroll(expected,"Payroll input")
analysis_input=analyse_input_payroll(expected,overtime_limit,overtime_hours_limit,bonus_limit)
before_items=pd.concat([attention_only(check_input),attention_only(analysis_input)],ignore_index=True)

employee_count=len(expected)
before_fails=int(before_items["Status"].eq("Fail").sum()) if not before_items.empty else 0
before_warn=int(before_items["Status"].eq("Warning").sum()) if not before_items.empty else 0
before_attention=len(before_items)

c1,c2,c3,c4=st.columns(4)
c1.metric("Employees",employee_count)
c2.metric("Before-payroll issues",before_attention)
c3.metric("Warnings",before_warn)
c4.metric("Fails",before_fails)

if before_attention==0:
    st.success("Before payroll: no input exceptions detected. You can proceed to your payroll system.")
else:
    st.markdown(f'<div class="exception-head"><b>{before_attention} item(s) need attention before payroll.</b> You do not need to inspect every employee — start with these exceptions.</div>',unsafe_allow_html=True)
    st.dataframe(before_items[[c for c in ["Status","Employee ID","Employee Name","Check / Issue","Detail"] if c in before_items.columns]],use_container_width=True,hide_index=True)

with st.expander("Show full before-payroll validation", expanded=False):
    st.dataframe(check_input,use_container_width=True,hide_index=True)
    st.dataframe(analysis_input,use_container_width=True,hide_index=True)

st.markdown('<div class="after-box"><div class="kicker">After payroll</div><div class="bigline">3. Upload the pre-release output</div><div class="muted">Once your payroll system has calculated payroll, upload its pre-release output in Excel, CSV, Word, text-based PDF, JSON/XML or another supported structured export. CARRR will preprocess and reconcile it automatically.</div></div>',unsafe_allow_html=True)
output_file=st.file_uploader("Pre-Release Payroll Output",type=["csv","tsv","txt","xlsx","xls","docx","pdf","json","xml","png","jpg","jpeg"],key="carrr_output")

if not output_file:
    st.info("Before-payroll checks are complete. Run payroll in your payroll system, then return here and upload the Pre-Release Payroll Output.")
    st.stop()

try:
    processed_raw=clean_columns(read_upload(output_file))
except Exception as exc:
    lower_name = output_file.name.lower()
    if lower_name.endswith((".pdf", ".png", ".jpg", ".jpeg")):
        verified_scan = render_ocr_verification(output_file, "output_ocr")
        if verified_scan is not None:
            processed_raw = clean_columns(verified_scan)
        else:
            st.stop()
    else:
        info=inspect_unstructured_source(output_file)
        st.warning("CARRR received the pre-release source but cannot safely structure it yet.")
        st.write(f"**Source:** {info.get('type','Unknown')}")
        st.write(f"**Status:** {info.get('status','Verification required')}")
        st.write(info.get("detail") or str(exc))
        if info.get("preview"):
            with st.expander("Extracted text preview", expanded=False):
                st.text(info["preview"])
        st.stop()

st.markdown("#### CARRR Auto Detect — Pre-Release Output")
processed, output_detect = auto_detect_and_prepare(processed_raw, output_file.name, "v25_output", role="output")
if processed is None:
    st.info("CARRR only needs confirmation for the unclear fields above before reconciliation can continue.")
    st.stop()

# v2.10 Source Validation — make extraction/population integrity visible BEFORE reconciliation.
source_validation = source_validation_summary(
    expected,
    processed,
    output_file.name,
    raw_source=processed_raw,
)

st.markdown("### Source Validation")
sv1,sv2,sv3,sv4,sv5=st.columns(5)
sv1.metric("Expected employees",source_validation["Expected Employees"])
sv2.metric("Extracted employees",source_validation["Extracted Employees"])
sv3.metric("Missing IDs",source_validation["Missing IDs"])
sv4.metric("Unexpected IDs",source_validation["Unexpected IDs"])
sv5.metric("Duplicate output IDs",source_validation["Duplicate Output ID Rows"])

source_population_aligned=bool(source_validation["Population Aligned"])

if source_population_aligned:
    st.success("✓ SOURCE READY — employee population is aligned and no duplicate/blank Employee IDs were detected.")
else:
    st.warning(
        "⚠ SOURCE REVIEW REQUIRED — the extracted pre-release population does not fully align with the expected payroll. "
        "CARRR will show the exact missing/unexpected employees during reconciliation instead of continuing silently."
    )

if source_validation.get("PDF Pages") is not None:
    pages=source_validation.get("PDF Pages")
    tables=source_validation.get("PDF Tables Combined")
    pages_with_tables=source_validation.get("PDF Pages With Tables")
    st.caption(
        f"PDF extraction check: {pages} page(s) detected; {tables} compatible payroll table(s) combined"
        + (f" across {pages_with_tables} page(s)." if pages_with_tables is not None else ".")
    )

with st.expander("Show source-validation details", expanded=not source_population_aligned):
    st.dataframe(source_validation_table(source_validation),use_container_width=True,hide_index=True)
    if source_validation["Missing ID List"]:
        st.write("**Missing Employee IDs:** "+", ".join(source_validation["Missing ID List"]))
    if source_validation["Unexpected ID List"]:
        st.write("**Unexpected Employee IDs:** "+", ".join(source_validation["Unexpected ID List"]))

common_fields=[f for f in RECON_FIELDS if f in expected.columns and f in processed.columns]
if not common_fields:
    st.error("The two files do not share a supported payroll value such as Gross Pay, PAYE, Pension or Net Pay.")
    st.stop()

with st.expander("Reconciliation settings",expanded=False):
    s1,s2=st.columns(2)
    fields=s1.multiselect("Fields to compare",common_fields,default=common_fields)
    tolerance=s2.number_input("Allowed difference (£)",min_value=0.0,value=0.01,step=0.01)

if not fields:
    st.warning("Select at least one field to compare.")
    st.stop()

errors=validate_reconciliation_columns(expected,"Payroll input",fields)+validate_reconciliation_columns(processed,"Pre-release output",fields)
if errors:
    st.error("CARRR cannot reconcile until these file-structure issues are fixed:")
    for e in errors: st.write("- "+e)
    st.stop()

check_output=check_payroll(processed,"Pre-release output")
reconciliation=reconcile(expected,processed,fields,tolerance)
controls=control_totals(reconciliation,fields)
recon_items=reconciliation_review_items(reconciliation)
after_items=pd.concat([attention_only(check_output),attention_only(recon_items)],ignore_index=True)
all_items=pd.concat([before_items,after_items],ignore_index=True)

exceptions_count=int(reconciliation["Overall Status"].eq("Exception").sum())
matched_count=len(reconciliation)-exceptions_count
net_variance=float(reconciliation["Net Pay Difference"].sum()) if "Net Pay Difference" in reconciliation.columns else 0.0

st.markdown("### CARRR result")
r1,r2,r3,r4=st.columns(4)
r1.metric("Employees checked",len(reconciliation))
r2.metric("Matched / OK",matched_count)
r3.metric("Employees with exceptions",exceptions_count)
r4.metric("Net pay variance",f"£{net_variance:,.2f}")

# Split true payroll mismatches from threshold/master-data warnings so users can fix the important records first.
mismatch_rows = reconciliation[reconciliation["Overall Status"].eq("Exception")].copy()
if not mismatch_rows.empty:
    def mismatch_type(row):
        if row.get("Record Status") == "Missing from processed payroll":
            return "Missing employee/output"
        if row.get("Record Status") == "Unexpected employee in processed payroll":
            return "Unexpected employee/payment"
        return "Payroll value mismatch"

    mismatch_summary = pd.DataFrame({
        "Employee ID": mismatch_rows["Employee ID"],
        "Employee Name": mismatch_rows["Employee Name"],
        "Mismatch Type": mismatch_rows.apply(mismatch_type, axis=1),
        "What changed": mismatch_rows["Exception Reason"],
        "Net Pay Difference": mismatch_rows["Net Pay Difference"] if "Net Pay Difference" in mismatch_rows.columns else 0.0,
    })
    mismatch_summary["Action"] = mismatch_summary["Mismatch Type"].map({
        "Missing employee/output": "Check payroll output / missing employee",
        "Unexpected employee/payment": "Verify employee and payment before release",
        "Payroll value mismatch": "Check source input and payroll calculation",
    })
else:
    mismatch_summary = pd.DataFrame(columns=["Employee ID","Employee Name","Mismatch Type","What changed","Net Pay Difference","Action"])

other_after_warnings = attention_only(check_output)

st.markdown("### Fix First — Payroll Mismatches")
if mismatch_summary.empty:
    st.success("No payroll value mismatches found.")
else:
    st.markdown(
        f'<div class="exception-head"><b>{len(mismatch_summary)} employee(s) have payroll mismatches.</b> '
        'These are the records to investigate first. They are separated from general warnings so you do not need to search through the full review queue.</div>',
        unsafe_allow_html=True,
    )
    show_mismatch = mismatch_summary.copy()
    if "Net Pay Difference" in show_mismatch.columns:
        show_mismatch["Net Pay Difference"] = show_mismatch["Net Pay Difference"].map(lambda x: f"£{x:,.2f}")
    st.dataframe(show_mismatch, use_container_width=True, hide_index=True)

if all_items.empty:
    st.success("No review exceptions detected.")
else:
    st.markdown(
        f'<div class="exception-head"><b>Review Centre: {len(all_items)} total review item(s).</b> '
        f'Payroll mismatches: {len(mismatch_summary)} employee(s). Before-payroll warnings: {len(before_items)} item(s). '
        'Use the separate sections below instead of searching one long table.</div>',
        unsafe_allow_html=True,
    )

# Build review queue from attention items only; no pass rows cluttering the screen.
review_queue=all_items.copy()
if review_queue.empty:
    review_queue=pd.DataFrame(columns=["Source","Status","Employee ID","Employee Name","Check / Issue","Detail"])
review_queue.insert(0,"Exception ID",[f"CARRR-{i+1:04d}" for i in range(len(review_queue))])
review_queue["Reviewed"]=False
review_queue["Review Status"]="Open"
review_queue["Reviewer Note"]=""

main_tab,mismatch_tab,warning_tab,detail_tab,review_tab,release_tab=st.tabs(["Overview","Fix First: Mismatches","Warnings","Reconciliation Detail","Review & Notes","Release Gate"])

with main_tab:
    st.markdown("#### What needs attention")
    o1,o2,o3,o4=st.columns(4)
    o1.metric("Payroll mismatches", len(mismatch_summary))
    o2.metric("Before-payroll warnings", len(before_items))
    o3.metric("Other output warnings", len(other_after_warnings))
    o4.metric("Source population", "Ready" if source_population_aligned else "Review")
    st.caption(
        "Start with Source Validation and Payroll mismatches. "
        "Warnings are secondary checks and can be reviewed afterwards."
    )
    with st.expander("Payroll control totals",expanded=False):
        shown=controls.copy()
        for col in ["Expected Total","Processed Total","Variance"]:
            if col in shown.columns: shown[col]=shown[col].map(lambda x:f"£{x:,.2f}")
        st.dataframe(shown,use_container_width=True,hide_index=True)

with mismatch_tab:
    st.markdown("#### Payroll mismatches — investigate these first")
    if mismatch_summary.empty:
        st.success("No payroll mismatches found.")
    else:
        mm=mismatch_summary.copy()
        mm["Net Pay Difference"] = mm["Net Pay Difference"].map(lambda x:f"£{x:,.2f}")
        st.dataframe(mm,use_container_width=True,hide_index=True)
        st.info("A mismatch shows that the pre-release payroll output differs from the expected/input data. Check the source input and the payroll-system calculation before release.")

with warning_tab:
    st.markdown("#### Warnings — review after payroll mismatches")
    st.caption("Warnings can include high overtime, missing master data, leaver payments or other threshold checks. They are separated from true reconciliation mismatches.")
    if before_items.empty and other_after_warnings.empty:
        st.success("No warnings found.")
    else:
        if not before_items.empty:
            st.markdown(f"**Before-payroll warnings ({len(before_items)})**")
            st.dataframe(before_items[[c for c in ["Status","Employee ID","Employee Name","Check / Issue","Detail"] if c in before_items.columns]],use_container_width=True,hide_index=True)
        if not other_after_warnings.empty:
            st.markdown(f"**Pre-release output warnings ({len(other_after_warnings)})**")
            st.dataframe(other_after_warnings[[c for c in ["Status","Employee ID","Employee Name","Check / Issue","Detail"] if c in other_after_warnings.columns]],use_container_width=True,hide_index=True)

with detail_tab:
    st.caption("Use this detailed table only when you need to investigate an individual employee or field difference.")
    st.dataframe(reconciliation,use_container_width=True,hide_index=True)

with review_tab:
    st.markdown("#### Record review decisions and notes")
    if review_queue.empty:
        review_editor=review_queue.copy()
        st.success("No exceptions require human review.")
    else:
        review_editor=st.data_editor(
            review_queue,use_container_width=True,hide_index=True,
            disabled=["Exception ID","Source","Status","Employee ID","Employee Name","Check / Issue","Detail"],
            column_config={
                "Reviewed":st.column_config.CheckboxColumn("Reviewed"),
                "Review Status":st.column_config.SelectboxColumn("Outcome",options=["Open","Resolved","Accepted"],required=True),
                "Reviewer Note":st.column_config.TextColumn("Reviewer Note"),
            },key="carrr_v2_review")
        open_items=review_editor[(~review_editor["Reviewed"])|review_editor["Review Status"].eq("Open")]
        st.metric("Total review items still open",len(open_items))

with release_tab:
    st.markdown("#### CARRR Release Gate")
    st.caption(
        "CARRR calculates the release position from the review evidence first. "
        "Manual approval cannot override an unresolved payroll control."
    )

    if review_queue.empty:
        final_review=review_queue.copy()
    else:
        final_review=review_editor.copy()

    # A review item is resolved only when the reviewer has explicitly marked it
    # reviewed AND selected either Resolved or Accepted.
    if final_review.empty:
        open_items=final_review.copy()
        accepted_without_note=final_review.copy()
    else:
        reviewed_flag=final_review["Reviewed"].fillna(False).astype(bool)
        review_status=final_review["Review Status"].fillna("Open").astype(str)
        open_items=final_review[(~reviewed_flag)|review_status.eq("Open")].copy()

        accepted_mask=reviewed_flag & review_status.eq("Accepted")
        note_text=final_review["Reviewer Note"].fillna("").astype(str).str.strip()
        accepted_without_note=final_review[accepted_mask & note_text.eq("")].copy()

    # Classify the unresolved controls for a clear Release Gate decision.
    if open_items.empty:
        open_recon=open_items.copy()
        open_fails=open_items.copy()
        open_warnings=open_items.copy()
    else:
        open_recon=open_items[open_items["Source"].astype(str).eq("RECONCILE")].copy()
        open_fails=open_items[
            open_items["Status"].astype(str).eq("Fail")
            & ~open_items["Source"].astype(str).eq("RECONCILE")
        ].copy()
        open_warnings=open_items[open_items["Status"].astype(str).eq("Warning")].copy()

    gate_blocker_count=(
        len(open_recon)
        + len(open_fails)
        + len(open_warnings)
        + len(accepted_without_note)
    )
    controls_clear=(gate_blocker_count == 0)

    st.markdown("##### Automated control decision")
    g1,g2,g3,g4=st.columns(4)
    g1.metric("Open mismatches",len(open_recon))
    g2.metric("Open fails",len(open_fails))
    g3.metric("Open warnings",len(open_warnings))
    g4.metric("Accepted without note",len(accepted_without_note))

    if controls_clear:
        st.markdown(
            '<div class="gate-auto-clear">✓ CARRR controls are clear. '
            'The payroll is eligible for final reviewer authorization.</div>',
            unsafe_allow_html=True,
        )
    else:
        st.markdown(
            f'<div class="gate-auto-hold">⚠ CARRR has {gate_blocker_count} release blocker(s). '
            'Payroll remains on hold until they are reviewed and resolved or formally accepted with a note.</div>',
            unsafe_allow_html=True,
        )

    # Build a compact blocker table for the reviewer and the exported audit file.
    blocker_frames=[]
    if not open_recon.empty:
        b=open_recon.copy()
        b.insert(0,"Blocker Type","Unresolved payroll mismatch")
        blocker_frames.append(b)
    if not open_fails.empty:
        b=open_fails.copy()
        b.insert(0,"Blocker Type","Unresolved fail")
        blocker_frames.append(b)
    if not open_warnings.empty:
        b=open_warnings.copy()
        b.insert(0,"Blocker Type","Unresolved warning")
        blocker_frames.append(b)
    if not accepted_without_note.empty:
        b=accepted_without_note.copy()
        b.insert(0,"Blocker Type","Accepted exception needs reviewer note")
        blocker_frames.append(b)

    if blocker_frames:
        blockers_df=pd.concat(blocker_frames,ignore_index=True)
        blocker_cols=[
            c for c in [
                "Blocker Type","Exception ID","Source","Status","Employee ID",
                "Employee Name","Check / Issue","Detail","Review Status","Reviewer Note"
            ] if c in blockers_df.columns
        ]
        blockers_df=blockers_df[blocker_cols]
        with st.expander("Show release blockers",expanded=True):
            st.dataframe(blockers_df,use_container_width=True,hide_index=True)
    else:
        blockers_df=pd.DataFrame(columns=[
            "Blocker Type","Exception ID","Source","Status","Employee ID",
            "Employee Name","Check / Issue","Detail","Review Status","Reviewer Note"
        ])
        st.success("No unresolved release blockers.")

    st.divider()
    st.markdown("##### Final authorization")

    p1,p2=st.columns(2)
    prepared_by=p1.text_input("Prepared by",key="release_prepared_by")
    reviewed_by=p2.text_input("Reviewed by",key="release_reviewed_by")
    review_date=st.date_input("Review date",value=date.today(),key="release_review_date")

    approval=st.selectbox(
        "Approval status",
        ["Pending","Approved","Hold"],
        help="Approved can only produce READY TO RELEASE when CARRR's automated controls are clear.",
        key="release_approval",
    )

    c1=st.checkbox("Payroll inputs checked",key="release_inputs_checked")
    c2=st.checkbox("All exceptions and warnings reviewed",key="release_exceptions_checked")
    c3=st.checkbox("Reconciliation and control totals reviewed",key="release_recon_checked")
    c4=st.checkbox(
        "I confirm the pre-release payroll is approved for release",
        key="release_final_confirmation",
    )

    names_complete=bool(prepared_by.strip()) and bool(reviewed_by.strip())
    checklist_complete=bool(c1 and c2 and c3 and c4)
    authorization_complete=(approval=="Approved" and names_complete and checklist_complete)

    ready=bool(controls_clear and authorization_complete)

    reasons=[]
    if not controls_clear:
        if len(open_recon): reasons.append(f"{len(open_recon)} payroll mismatch(es) unresolved")
        if len(open_fails): reasons.append(f"{len(open_fails)} fail item(s) unresolved")
        if len(open_warnings): reasons.append(f"{len(open_warnings)} warning(s) unresolved")
        if len(accepted_without_note): reasons.append(f"{len(accepted_without_note)} accepted exception(s) missing reviewer note")
    if not names_complete:
        reasons.append("prepared/reviewed by details incomplete")
    if approval=="Hold":
        reasons.append("reviewer selected Hold")
    elif approval!="Approved":
        reasons.append("final approval is pending")
    if not checklist_complete:
        reasons.append("release checklist incomplete")

    if ready:
        st.markdown('<div class="release-ready">✓ READY TO RELEASE</div>',unsafe_allow_html=True)
        decision="READY TO RELEASE"
        st.success(
            "All CARRR controls are clear, the review evidence is complete, "
            "and final authorization has been recorded."
        )
    else:
        st.markdown('<div class="release-hold">HOLD PAYROLL</div>',unsafe_allow_html=True)
        decision="HOLD PAYROLL"
        if reasons:
            st.markdown("**Why payroll is on hold:**")
            for reason in reasons:
                st.write(f"- {reason}")

    resolved_count=0
    accepted_count=0
    if not final_review.empty:
        status_text=final_review["Review Status"].fillna("Open").astype(str)
        reviewed_text=final_review["Reviewed"].fillna(False).astype(bool)
        resolved_count=int((reviewed_text & status_text.eq("Resolved")).sum())
        accepted_count=int((reviewed_text & status_text.eq("Accepted")).sum())

    summary_df=pd.DataFrame([
        {"Metric":"Payroll Period","Value":payroll_period or "Not specified"},
        {"Metric":"Employees Checked","Value":len(reconciliation)},
        {"Metric":"Matched / OK","Value":matched_count},
        {"Metric":"Reconciliation Exceptions","Value":exceptions_count},
        {"Metric":"Open Payroll Mismatches","Value":len(open_recon)},
        {"Metric":"Open Fails","Value":len(open_fails)},
        {"Metric":"Open Warnings","Value":len(open_warnings)},
        {"Metric":"Accepted Without Note","Value":len(accepted_without_note)},
        {"Metric":"Resolved Review Items","Value":resolved_count},
        {"Metric":"Accepted Review Items","Value":accepted_count},
        {"Metric":"Net Pay Variance","Value":round(net_variance,2)},
        {"Metric":"Source Expected Employees","Value":source_validation["Expected Employees"]},
        {"Metric":"Source Extracted Employees","Value":source_validation["Extracted Employees"]},
        {"Metric":"Source Missing IDs","Value":source_validation["Missing IDs"]},
        {"Metric":"Source Unexpected IDs","Value":source_validation["Unexpected IDs"]},
        {"Metric":"Source Duplicate Output IDs","Value":source_validation["Duplicate Output ID Rows"]},
        {"Metric":"Source Population Aligned","Value":"Yes" if source_population_aligned else "No"},
        {"Metric":"CARRR Controls Clear","Value":"Yes" if controls_clear else "No"},
        {"Metric":"Final Decision","Value":decision},
    ])

    release_df=pd.DataFrame([{
        "Payroll Period":payroll_period,
        "Prepared By":prepared_by,
        "Reviewed By":reviewed_by,
        "Review Date":str(review_date),
        "Approval Status":approval,
        "Inputs Checked":c1,
        "Exceptions Reviewed":c2,
        "Reconciliation Reviewed":c3,
        "Final Confirmation":c4,
        "Open Payroll Mismatches":len(open_recon),
        "Open Fails":len(open_fails),
        "Open Warnings":len(open_warnings),
        "Accepted Without Note":len(accepted_without_note),
        "Source Expected Employees":source_validation["Expected Employees"],
        "Source Extracted Employees":source_validation["Extracted Employees"],
        "Source Missing IDs":source_validation["Missing IDs"],
        "Source Unexpected IDs":source_validation["Unexpected IDs"],
        "Source Population Aligned":"Yes" if source_population_aligned else "No",
        "CARRR Controls Clear":"Yes" if controls_clear else "No",
        "Decision":decision,
        "Hold Reasons":"; ".join(reasons) if reasons else "",
    }])

    report=to_excel_bytes(
        summary_df,check_input,analysis_input,controls,reconciliation,
        final_review,release_df,blockers_df
    )
    fp=(payroll_period.strip().replace(" ","_") if payroll_period.strip() else "payroll")
    st.download_button(
        "Download CARRR audit report",
        report,
        file_name=f"CARRR_{fp}_audit_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
